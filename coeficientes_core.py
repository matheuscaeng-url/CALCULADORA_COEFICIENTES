"""Processo 3: tendências de tubos, conexões e eletrocalhas SINAPI.

O ajuste é feito por recurso, dentro de uma única competência e família.
Tês tubulares são tratados pelas três bocas. A estimativa não representa uma
nova aferição SINAPI e nunca cria um código oficial inexistente.
"""
from __future__ import annotations

import bisect
import gzip
import hashlib
import io
import json
import math
import re
import unicodedata
import zipfile
from collections import defaultdict
from datetime import datetime
from fractions import Fraction
from pathlib import Path

from dependencias_runtime import configure_dependencies
from justificativas import build_justification

configure_dependencies()

VERSION = "1.5.22"
AUTOMATIC_SCOPE = "Tubos, conexões e eletrocalhas"
SCOPE_REVISION = 2
OFFICIAL_URL = "https://www.caixa.gov.br/site/Paginas/downloads.aspx#categoria_888"


class AnalysisError(ValueError):
    pass


def norm(value):
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().upper()
    text = re.sub(r"AR[ -]+CONDICIONADO", "AR CONDICIONADO", text)
    return re.sub(r"\s+", " ", text).strip()


def family_signature_text(value):
    """Normaliza diferenças meramente tipográficas usadas no agrupamento de famílias.

    Não altera a descrição oficial armazenada. Serve apenas para que variações como
    ``AR-CONDICIONADO``/``AR CONDICIONADO``, ``DN 25MM``/``DN 25 MM`` e
    pequenos espaçamentos de pontuação não criem famílias artificiais distintas.
    """
    text = norm(value)
    # Hífens internos usados só como separador textual não devem quebrar a família.
    text = re.sub(r"(?<=\w)-(?=\w)", " ", text)
    # O marcador dimensional deve ter uma forma canônica, independentemente do
    # espaço existente no texto oficial entre DN, número e unidade.
    text = re.sub(r"\bDN\s*\{VALOR\}", "DN {VALOR}", text)
    text = re.sub(
        r"\{VALOR\}\s*(MM3|MM2|MM|CM3|CM2|CM|M3|M2|M|POL(?:EGADAS?)?|BTU(?:S)?/?H|KW|KVA|W|VA|V|A|HZ|KG|G|T|L|ML|BAR|MCA|PA|MPA|RPM|CV|HP)",
        r"{VALOR} \1", text)
    text = re.sub(r"\s*,\s*", ", ", text)
    text = re.sub(r"\s*;\s*", "; ", text)
    text = re.sub(r"\s*-\s*", " - ", text)
    return re.sub(r"\s+", " ", text).strip(" ,.-")


def number(value):
    if isinstance(value, bool) or value is None or str(value).strip() in ("", "-", "–"):
        raise ValueError("Número ausente")
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        s = str(value).strip().replace("\xa0", "").replace(" ", "")
        if "," in s:
            s = s.replace(".", "").replace(",", ".")
        result = float(s)
    if not math.isfinite(result):
        raise ValueError("Número não finito")
    return result


def code_text(value):
    """Normaliza códigos SINAPI vindos do Excel ou digitados pelo usuário.

    A referência pode entregar o código como número, texto ou número com
    representação decimal (ex.: 94464.0). Para consulta, essas formas devem
    apontar para o mesmo código oficial.
    """
    if isinstance(value, bool) or value is None:
        return ""
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            return ""
        if float(value).is_integer():
            return str(int(value))
        return str(value).strip()
    text = str(value).strip().replace("\xa0", "")
    # Excel/CSV pode transformar o código inteiro em texto decimal.
    if re.fullmatch(r"[+-]?\d+[.,]0+", text):
        try:
            return str(int(float(text.replace(",", "."))))
        except ValueError:
            pass
    # Remove apóstrofo usado pelo Excel para forçar célula como texto.
    if text.startswith("'"):
        text = text[1:].strip()
    return text


def resolve_composition_code(compositions, value):
    """Retorna a chave real da composição, aceitando representações equivalentes."""
    requested = code_text(value)
    if requested in compositions:
        return requested
    if not requested:
        return None
    # Normaliza também as chaves de bases antigas/importadas por versões anteriores.
    matches = [key for key in compositions if code_text(key) == requested]
    if len(matches) == 1:
        return matches[0]
    return None


def load_database(path):
    with gzip.open(path, "rt", encoding="utf-8") as stream:
        data = json.load(stream)
    if data.get("schema") != 2 or not data.get("compositions"):
        raise AnalysisError("Base anterior ou incompatível. Importe a referência SINAPI novamente.")
    return data


def save_database(data, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    with gzip.open(temp, "wt", encoding="utf-8") as stream:
        json.dump(data, stream, ensure_ascii=False, separators=(",", ":"))
    temp.replace(path)


def import_sinapi(path, progress=lambda message: None, include_all=False):
    """Lê o escopo automático; include_all preserva todas as CPUs para a aba manual."""
    from openpyxl import load_workbook
    path = Path(path)
    progress("Lendo o arquivo de referência…")
    source_bytes = path.read_bytes()
    source_name = path.name
    if path.suffix.lower() == ".zip":
        try:
            with zipfile.ZipFile(io.BytesIO(source_bytes)) as archive:
                candidates = [n for n in archive.namelist() if n.lower().endswith(".xlsx")
                              and "REFERENCIA" in norm(Path(n).name)]
                if len(candidates) != 1:
                    raise AnalysisError("Extraia o ZIP e selecione a planilha SINAPI_Referência do mês desejado.")
                item = archive.getinfo(candidates[0])
                if item.file_size > 250_000_000:
                    raise AnalysisError("Planilha muito grande. Extraia o ZIP e selecione o XLSX.")
                source_name = Path(candidates[0]).name
                source_bytes = archive.read(candidates[0])
        except zipfile.BadZipFile as exc:
            raise AnalysisError("O ZIP está incompleto ou danificado. Baixe o arquivo novamente.") from exc
    try:
        workbook = load_workbook(io.BytesIO(source_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise AnalysisError("Não foi possível abrir o Excel. Use o XLSX completo, extraído do ZIP da CAIXA.") from exc
    try:
        sheet = next((s for s in workbook if norm(s.title) in ("ANALITICO", "ANALITICO DE COMPOSICOES")), None)
        if sheet is None:
            raise AnalysisError("Não encontrei a aba Analítico completa. Selecione SINAPI_Referência, com os coeficientes; uma tabela de preços não basta.")
        rows = sheet.iter_rows(values_only=True)
        period, emission, columns, header_row = None, None, None, None
        for index, row in enumerate(rows, 1):
            joined = " ".join(str(c or "") for c in row)
            if "MES DE REFERENCIA" in norm(joined):
                match = re.search(r"\b(0?[1-9]|1[0-2])/(20\d{2})\b", joined)
                if match:
                    period = f"{int(match[1]):02d}/{match[2]}"
            if "DATA DE EMISSAO" in norm(joined):
                match = re.search(r"\b\d{2}/\d{2}/\d{4}\b", joined)
                emission = match[0] if match else None
            normalized = [norm(c) for c in row]
            wanted = {"code": "CODIGO DA COMPOSICAO", "kind": "TIPO ITEM", "item": "CODIGO DO ITEM",
                      "description": "DESCRICAO", "unit": "UNIDADE", "coefficient": "COEFICIENTE"}
            if all(label in normalized for label in wanted.values()):
                columns = {key: normalized.index(label) for key, label in wanted.items()}
                for key, label in (("group", "GRUPO"), ("status", "SITUACAO")):
                    if label in normalized:
                        columns[key] = normalized.index(label)
                header_row = index
                break
            if index >= 100:
                break
        if not columns:
            raise AnalysisError("Cabeçalhos analíticos não reconhecidos: código da composição, tipo item, código do item, descrição, unidade e coeficiente.")
        if not period:
            raise AnalysisError("Não identifiquei o mês de referência no cabeçalho. A data AF da descrição não é a competência.")
        compositions, pending, missing = {}, defaultdict(list), 0
        def get(row, key):
            i = columns.get(key)
            return row[i] if i is not None and i < len(row) else None
        for index, row in enumerate(rows, header_row + 1):
            parent = code_text(get(row, "code"))
            description = str(get(row, "description") or "").strip()
            if not parent or not description:
                continue
            item = code_text(get(row, "item"))
            kind = norm(get(row, "kind"))
            if not item and not kind:
                if parent in compositions:
                    raise AnalysisError(f"Composição {parent} aparece duas vezes no relatório; confira a referência.")
                compositions[parent] = {"code": parent, "description": description,
                    "unit": str(get(row, "unit") or "").strip(), "group": str(get(row, "group") or "").strip(),
                    "status": str(get(row, "status") or ""), "row": index, "components": pending[parent]}
            elif item:
                try:
                    coefficient = number(get(row, "coefficient"))
                    if coefficient < 0:
                        raise ValueError()
                except (ValueError, TypeError):
                    coefficient = None
                    missing += 1
                pending[parent].append({"code": item, "kind": kind, "description": description,
                    "unit": str(get(row, "unit") or "").strip(), "coefficient": coefficient,
                    "status": str(get(row, "status") or ""), "row": index})
            if index % 10000 == 0:
                progress(f"Lendo coeficientes: {index:,} linhas…")
        if len(compositions) < 3:
            raise AnalysisError("O relatório não contém composições suficientes para identificar tendências.")
        orphan = sorted(set(pending) - set(compositions))
        if orphan:
            raise AnalysisError(f"Há itens sem composição principal: {', '.join(orphan[:5])}. Confira se a planilha está completa.")
        raw_count = len(compositions)
        tray_materials = {component["code"] for item in compositions.values()
                         if is_cable_tray(item["description"])
                         for component in item["components"] if component["kind"] == "INSUMO"}
        if not include_all:
            compositions = {code: item for code, item in compositions.items() if is_automatic_scope(item["description"])}
        if len(compositions) < 3:
            raise AnalysisError("Não encontrei composições suficientes de tubos, conexões e eletrocalhas na referência selecionada.")

        supplies, price_sheet = {}, None
        price_ws = next((s for s in workbook if norm(s.title) == "ISD"), None)
        if price_ws is not None:
            price_sheet = price_ws.title
            price_rows = price_ws.iter_rows(values_only=True)
            price_columns = None
            for price_row_number, price_row in enumerate(price_rows, 1):
                labels = [norm(c) for c in price_row]
                required = ["CODIGO DO INSUMO", "DESCRICAO DO INSUMO", "UNIDADE", "PE", "SP"]
                if all(label in labels for label in required):
                    price_columns = {label: labels.index(label) for label in required}
                    for optional in ("CLASSIFICACAO", "ORIGEM DE PRECO"):
                        if optional in labels:
                            price_columns[optional] = labels.index(optional)
                    break
                if price_row_number >= 100:
                    break
            if price_columns:
                def price_get(row, label):
                    i = price_columns.get(label)
                    return row[i] if i is not None and i < len(row) else None
                for price_row_number, price_row in enumerate(price_rows, price_row_number + 1):
                    code = code_text(price_get(price_row, "CODIGO DO INSUMO"))
                    description = str(price_get(price_row, "DESCRICAO DO INSUMO") or "").strip()
                    classification = norm(price_get(price_row, "CLASSIFICACAO"))
                    if (not code or not description or classification != "MATERIAL"
                            or not (is_automatic_scope(description) or code in tray_materials)):
                        continue
                    def optional_price(state):
                        try:
                            return number(price_get(price_row, state))
                        except (ValueError, TypeError):
                            return None
                    supplies[code] = {"code": code, "description": description,
                        "unit": str(price_get(price_row, "UNIDADE") or "").strip(),
                        "origin": str(price_get(price_row, "ORIGEM DE PRECO") or "").strip(),
                        "PE": optional_price("PE"), "SP": optional_price("SP"),
                        "row": price_row_number, "sheet": price_ws.title}
        return {"schema": 2, "meta": {"period": period, "emission": emission, "filename": source_name,
                "sha256": hashlib.sha256(source_bytes).hexdigest(), "sheet": sheet.title,
                "price_sheet": price_sheet, "imported_at": datetime.now().isoformat(timespec="seconds"),
                "missing_coefficients": missing, "source_compositions": raw_count,
                "filtered_compositions": len(compositions), "filtered_supplies": len(supplies),
                "scope": "Todas as composições (seleção manual)" if include_all else AUTOMATIC_SCOPE,
                "scope_revision": SCOPE_REVISION, "source_url": OFFICIAL_URL},
                "compositions": compositions, "supplies": supplies}
    finally:
        workbook.close()


# Dimensões e escopo técnico -------------------------------------------------
# Tubulações e eletrocalhas têm reconhecedores distintos: um tê elétrico
# não representa as três bocas de um tê de tubulação.
MEASURE = re.compile(r"(?<![\w.,])(?P<value>\d+(?:[.,]\d+)?)\s*(?P<unit>MM2|MM²|CM2|CM²|BTU(?:S)?\s*/?\s*H|KVA|KW|MM|CM|M|A)\b", re.I)
INCH = re.compile(r'(?<![\w.,])(?P<value>\d+(?:\s+\d+/\d+|/\d+|[.,]\d+)?)\s*(?:"|POL(?:EGADAS?)?\b)', re.I)

PIPE_MATERIAL = re.compile(
    r"\b(PVC|CPVC|PPR|PEX|PEAD|POLIETILENO|POLIPROPILENO|COBRE|BRONZE|LATAO|"
    r"FERRO(?: FUNDIDO| GALVANIZADO)?|ACO(?: GALVANIZADO| CARBONO| INOXIDAVEL)?|FIBROCIMENTO)\b")
FITTING = re.compile(
    r"\b(TE|JOELHO|COTOVELO|CURVA|LUVA|UNIAO|ADAPTADOR|BUCHA|NIPLE|CAP|TAMPAO|"
    r"JUNCAO|CRUZETA|FLANGE|REDUCAO|CONECTOR|CONEXAO|COLAR|SELIM)\b")
SCOPE_EXCLUSION = re.compile(
    r"\b(ELETRODUTO|ELETROCALHA|BANDEJA|LEITO PARA CABOS|CABO ELETRICO|CONDULETE|"
    r"DUTO DE AR|DUTO PARA AR|CHAPA GALVANIZADA|CALHA DE PISO|CANALETA ELETRICA)\b")

_FRACTIONS = {"½": "1/2", "¼": "1/4", "¾": "3/4", "⅛": "1/8", "⅜": "3/8", "⅝": "5/8", "⅞": "7/8"}
_DIM_NUMBER = r"(?:\d+\s+\d+/\d+|\d+/\d+|\d+(?:[.,]\d+)?)"
_DIM_UNIT = r'(?:MM|CM|M|POL(?:EGADAS?)?|\")'
_DIM_TOKEN = rf"{_DIM_NUMBER}\s*{_DIM_UNIT}?"
_DN_SEQUENCE = re.compile(
    rf"\b(?:DN|DIAMETRO(?:\s+NOMINAL)?(?:\s+DE)?)\s*=?\s*(?P<seq>{_DIM_TOKEN}(?:\s*X\s*{_DIM_TOKEN}){{0,2}})", re.I)
_BARE_SEQUENCE = re.compile(rf"(?P<seq>{_DIM_TOKEN}\s*X\s*{_DIM_TOKEN}(?:\s*X\s*{_DIM_TOKEN})?)", re.I)
_DIM_PART = re.compile(rf"(?P<number>{_DIM_NUMBER})\s*(?P<unit>{_DIM_UNIT})?", re.I)

# Equivalências nominais usuais de tubulações. Elas não representam conversão
# geométrica por 25,4: DN 80 e NPS 3\", por exemplo, nomeiam o mesmo tubo.
_DN_TO_INCH = {
    6: 0.125, 8: 0.25, 10: 0.375, 15: 0.5, 20: 0.75, 25: 1.0,
    32: 1.25, 40: 1.5, 50: 2.0, 65: 2.5, 80: 3.0, 100: 4.0,
    125: 5.0, 150: 6.0, 200: 8.0, 250: 10.0, 300: 12.0,
    350: 14.0, 400: 16.0, 450: 18.0, 500: 20.0, 550: 22.0,
    600: 24.0, 650: 26.0, 700: 28.0, 750: 30.0, 800: 32.0,
    900: 36.0, 1000: 40.0, 1100: 44.0, 1200: 48.0,
    1300: 52.0, 1400: 56.0, 1500: 60.0, 1600: 64.0,
    1700: 68.0, 1800: 72.0, 1900: 76.0, 2000: 80.0, 2100: 84.0,
}


def is_tube_or_fitting(description):
    """True apenas para serviços/insumos cujo objeto principal é tubo ou conexão."""
    text = norm(description)
    if not text or SCOPE_EXCLUSION.search(text):
        return False
    pipe_context = bool(PIPE_MATERIAL.search(text) or re.search(
        r"\b(DN|TUBULACAO|REDE DE AGUA|REDE DE ESGOTO|ESGOTO PREDIAL|AGUA PLUVIAL|"
        r"HIDRAULICA PREDIAL|GAS COMBUSTIVEL|HIDRANTE|SPRINKLER)\b", text))
    if not pipe_context:
        return False
    # "Tubo" deve ser o objeto executado, não mera menção em pintura, ensaio ou suporte.
    if re.search(r"^(?:FORNECIMENTO E )?(?:ASSENTAMENTO E FORNECIMENTO DE |FORNECIMENTO E ASSENTAMENTO DE |"
                 r"ASSENTAMENTO DE |MONTAGEM DE |INSTALACAO DE )?TUBOS?\b", text):
        return True
    if re.search(r"\b(?:FORNECIMENTO E ASSENTAMENTO|ASSENTAMENTO E FORNECIMENTO|MONTAGEM) DE TUBOS?\b", text):
        return True
    match = FITTING.search(text)
    return bool(match and match.start() < 95)


def is_pipe_tee(description):
    text = norm(description)
    return is_tube_or_fitting(description) and bool(re.search(r"\bTE\b", text))


def is_cable_tray(description):
    """O objeto executado é a eletrocalha, sua conexão ou seu suporte.

    Uma mera menção, como cabo instalado em eletrocalha, não amplia o escopo.
    """
    text = norm(description)
    if not re.search(r"\bELETROCALHAS?\b", text):
        return False
    return bool(re.match(
        r"^(?:(?:FORNECIMENTO E )?INSTALACAO DE |FORNECIMENTO DE )?"
        r"(?:ELETROCALHAS?|COTOVELOS?|CURVAS?|TE|EMENDAS?|REDUCAO|SUPORTES?|"
        r"TALAS?|TAMPAS?|TERMINAIS?|CRUZETAS?|JUNCAO|ACOPLAMENTOS?|DERIVACAO|ADAPTADORES?|SAIDAS?)\b", text))


def is_automatic_scope(description):
    return is_tube_or_fitting(description) or is_cable_tray(description)


def _normalized_dimensions_text(description):
    text = str(description or "")
    for symbol, replacement in _FRACTIONS.items():
        text = text.replace(symbol, replacement)
    text = re.sub(r"(\d+)[.-](\d+/\d+)(?=\s*(?:\"|POL))", r"\1 \2", text, flags=re.I)
    # Há descrições históricas com uma vírgula antes do X dimensional.
    text = re.sub(r",\s*X\s*(?=\d)", " X ", text, flags=re.I)
    return norm(text)


def _parse_dimension_value(value):
    value = value.strip().replace(",", ".")
    if " " in value and "/" in value:
        whole, fraction = value.split(None, 1)
        numerator, denominator = fraction.split("/")
        return float(whole) + float(numerator) / float(denominator)
    if "/" in value:
        numerator, denominator = value.split("/")
        return float(numerator) / float(denominator)
    return float(value)


_DIMENSION_SUFFIX = re.compile(r'\s*(MM|CM|M|POL(?:EGADAS?)?|["″“”])\s*$', re.I)


def dimension_unit_suffix(value):
    """Unidade escrita no próprio campo, para conferir a unidade da coluna."""
    match = _DIMENSION_SUFFIX.search(str(value))
    if not match:
        return None
    unit = match[1].upper()
    return 'POL' if unit.startswith('POL') or unit in ('"', '″', '“', '”') else unit


def dimension_number(value):
    """Lê decimal, fração simples ou mista, com unidade opcional.

    A fração é uma notação do número: 3/4 POL continua sendo 0,75 POL.
    A conversão de unidades pertence ao cálculo, não à leitura do campo.
    """
    if isinstance(value, (int, float)):
        return number(value)
    text = str(value or "").strip()
    for symbol, replacement in _FRACTIONS.items():
        # Em 1½, o espaço separa a parte inteira da fração (não é 11/2).
        text = text.replace(symbol, ' ' + replacement)
    text = text.replace('⁄', '/').replace('∕', '/')
    text = _DIMENSION_SUFFIX.sub('', text).strip()
    text = re.sub(r'\s*/\s*', '/', text)
    text = re.sub(r"(\d+)[.-](\d+/\d+)", r"\1 \2", text)
    try:
        if '/' in text:
            match = re.fullmatch(r'([+-]?)(?:(\d+)\s+)?(\d+)/(\d+)', text)
            if not match:
                raise ValueError('Fração inválida')
            sign, whole, numerator, denominator = match.groups()
            result = int(whole or 0) + int(numerator) / int(denominator)
            if sign == '-':
                result = -result
        else:
            result = number(text)
    except (ValueError, ZeroDivisionError) as exc:
        raise ValueError(f"Dimensão inválida: {value}") from exc
    if not math.isfinite(result):
        raise ValueError("Dimensão não finita")
    return result


def _dimension_sequence(description):
    """Extrai uma sequência dimensional principal de uma conexão."""
    text = _normalized_dimensions_text(description)
    match = _DN_SEQUENCE.search(text) or _BARE_SEQUENCE.search(text)
    if not match:
        return None
    parts = list(_DIM_PART.finditer(match["seq"]))
    if not 1 <= len(parts) <= 3:
        return None
    dims = [_parse_dimension_value(part["number"]) for part in parts]
    units = [norm(part["unit"]).replace("POLEGADAS", "POL").replace("POLEGADA", "POL").replace('"', "POL")
             if part["unit"] else "" for part in parts]
    explicit = [bool(part["unit"]) for part in parts]
    # Em várias conexões roscáveis antigas a medida da rosca aparece sem o
    # símbolo de polegada. O material e o tipo de junta tornam a unidade clara.
    mixed_thread = bool(re.search(r"\b(LATAO|BRONZE|BUCHA|ROSCA|ROSQUEAD[AO])\b", text))
    if len(parts) > 1 and units[0] == "MM" and mixed_thread:
        units = [unit or ("POL" if index else "MM") for index, unit in enumerate(units)]
    else:
        inferred = next((unit for unit in reversed(units) if unit), "MM")
        units = [unit or inferred for unit in units]
    if any(value <= 0 for value in dims):
        return None
    return {"dims": dims, "units": units, "span": match.span(), "text": text,
            "reported_dimensions": len(dims), "inferred_units": [not flag for flag in explicit],
            "has_dn": bool(_DN_SEQUENCE.search(text) and _DN_SEQUENCE.search(text).span() == match.span())}


def cable_tray_geometry(description):
    """Extrai largura e altura; reduções têm duas larguras e uma altura.

    L × A de um insumo não são duas larguras; (L1 × L2) × A é redução.
    Suportes têm largura atendida e comprimento do perfil, com a folga entre
    ambos preservada no agrupamento e na entrada do alvo.
    """
    if not is_cable_tray(description):
        return None
    text = _normalized_dimensions_text(description)
    value = r"\d+(?:[.,]\d+)?"
    scalar = rf"(?P<value>{value})\s*(?P<unit>MM|CM)\b"
    height = re.search(rf"\bALTURA\s*(?:DE\s*)?{scalar}", text)
    width = re.search(rf"\bLARGURA\s*(?:DE\s*)?{scalar}", text)
    support = bool(re.match(r"SUPORTE\b", text))
    reduction = bool(re.match(r"REDUCAO\b", text))

    def reading(match):
        return number(match["value"]), match["unit"], (match.start("value"), match.end("unit"))

    if support:
        length = re.search(rf"\bCOMPRIMENTO\s*(?:DE\s*)?{scalar}", text)
        if not width or not length:
            return None
        readings = [reading(width), reading(length)]
        height_mm = None
        labels = ["Largura atendida", "Comprimento do perfil"]
        kind = "tray_support"
    elif width and height and not reduction:
        readings = [reading(width), reading(height)]
        height_mm = dimension_to_mm(*reading(height)[:2])
        labels, kind = ["Largura", "Altura"], "tray"
    elif height and not width and not reduction:
        # Talas e acessórios especificados somente pela altura também mudam
        # quando o usuário adapta a altura da eletrocalha.
        readings = [reading(height)]
        height_mm = dimension_to_mm(*readings[0][:2])
        labels, kind = ["Altura"], "tray_height_accessory"
    else:
        pair = re.search(
            rf"(?<![\w.,])(?P<a>{value})\s*(?P<ua>MM|CM)?\s*X\s*"
            rf"(?P<b>{value})\s*(?P<ub>MM|CM)\b", text)
        if not pair:
            return None
        unit_a, unit_b = pair["ua"] or pair["ub"], pair["ub"]
        first = (number(pair["a"]), unit_a, (pair.start("a"), pair.end("ua") if pair["ua"] else pair.end("a")))
        second = (number(pair["b"]), unit_b, (pair.start("b"), pair.end("ub")))
        if reduction:
            if not height:
                height = re.search(rf"\)?\s*X\s*{scalar}", text[pair.end():])
                height_offset = pair.end()
            else:
                height_offset = 0
            if not height:
                return None
            height_mm = dimension_to_mm(number(height["value"]), height["unit"])
            height_reading = (number(height["value"]), height["unit"],
                              (height.start("value") + height_offset, height.end("unit") + height_offset))
            readings = [first, second, height_reading]
            labels, kind = ["Largura maior", "Largura menor", "Altura"], "tray_reduction"
        else:
            # L × A são dois parâmetros, mesmo quando possuem o mesmo valor.
            height_mm = dimension_to_mm(second[0], second[1])
            readings = [first, second]
            labels, kind = ["Largura", "Altura"], "tray"
    dims, units, spans = map(list, zip(*readings))
    if any(d <= 0 for d in dims) or (height_mm is not None and height_mm <= 0):
        return None
    return {"dims": dims, "units": units, "spans": spans, "text": text,
            "reported_dimensions": len(dims), "inferred": False, "kind": kind,
            "axis_labels": labels, "height_mm": height_mm,
            "profile_offset_mm": (dimension_to_mm(dims[1], units[1]) - dimension_to_mm(dims[0], units[0])
                                  if support else None)}


def cable_tray_signature(description, geometry=None, *, align_gauge=False):
    geometry = geometry or cable_tray_geometry(description)
    if not geometry:
        return None
    text = geometry["text"]
    for index in reversed(range(len(geometry["spans"]))):
        start, end = geometry["spans"][index]
        text = text[:start] + f"{{D{index + 1}}}" + text[end:]
    text = re.sub(r"\b(LARGURA|ALTURA|COMPRIMENTO)\s+DE\s+", r"\1 ", text)
    if geometry["kind"] == "tray_reduction":
        text = re.sub(r"\bLARGURA\s+(?=\{D1\})", "", text)
    text = re.sub(r"\(\s*L\s*X\s*A\s*\)", "", text)
    text = re.sub(r"\b(\d+)O\b", r"\1 GRAUS", text)
    text = re.sub(r"\bAF[_ -]?\d{2}/\d{4}\b", "", text)
    # A CPU não fixa a bitola da chapa. Para alinhar coeficientes, aceita-se
    # a troca de bitola observada na própria família. Preços mantêm a bitola.
    if align_gauge:
        text = re.sub(r"\bESPESSURA\s*#\s*\d+|#\s*\d+", "", text)
    text = re.sub(r"(?<=\d)\s*(MM|CM)\b", r" \1", text)
    text = re.sub(r"\s*X\s*(?=\{D|\d)", " X ", text)
    text = re.sub(r"#\s+", "#", text)
    text = text.replace(",", " ")
    return re.sub(r"\s+", " ", text).strip(" .-")


def tee_geometry(description):
    """Retorna sempre as três bocas de um tê e a origem textual da inferência."""
    if not is_pipe_tee(description):
        return None
    parsed = _dimension_sequence(description)
    if not parsed:
        text = _normalized_dimensions_text(description)
        # Alguns tês antigos informam somente 4" sem o prefixo DN.
        inch = INCH.search(text)
        if not inch:
            return None
        dims, units, span, has_dn = [inch_number(inch["value"])], ["POL"], inch.span(), False
    else:
        text, dims, units, span, has_dn = (parsed["text"], list(parsed["dims"]),
            list(parsed["units"]), parsed["span"], parsed["has_dn"])
    original_count = len(dims)
    if original_count == 1:
        dims, units = dims * 3, units * 3
    elif original_count == 2:
        dims, units = [dims[0], dims[0], dims[1]], [units[0], units[0], units[1]]
    if any(value <= 0 for value in dims):
        return None
    return {"dims": dims, "units": units, "span": span, "text": text,
            "reported_dimensions": original_count, "inferred": original_count < 3,
            "has_dn": has_dn, "kind": "tee"}


def connection_geometry(description):
    """Geometria completa de conexão com duas ou três bocas/dimensões.

    Tês sempre têm três bocas. Nas demais conexões, todas as dimensões ligadas
    por X variam juntas; isso evita criar uma família diferente para cada lado
    fixo de uma redução.
    """
    if is_cable_tray(description):
        return None
    tee = tee_geometry(description)
    if tee:
        return tee
    text = norm(description)
    fitting = FITTING.search(text)
    if not fitting or fitting.start() >= 95 or re.search(r"^TUBOS?\b", text):
        return None
    parsed = _dimension_sequence(description)
    if not parsed or len(parsed["dims"]) < 2:
        return None
    parsed["inferred"] = any(parsed.pop("inferred_units", []))
    parsed["kind"] = "connection"
    return parsed


def dimension_to_mm(value, unit):
    factors = {"MM": 1.0, "CM": 10.0, "M": 1000.0, "POL": 25.4}
    if norm(unit) not in factors:
        raise AnalysisError(f"Unidade dimensional não suportada para tê: {unit}.")
    return number(value) * factors[norm(unit)]


def dimension_from_mm(value, unit):
    factors = {"MM": 1.0, "CM": 10.0, "M": 1000.0, "POL": 25.4}
    if norm(unit) not in factors:
        raise AnalysisError(f"Unidade dimensional não suportada: {unit}.")
    return number(value) / factors[norm(unit)]


def nominal_dimensions_equal(value_a, unit_a, value_b, unit_b):
    """Compara diâmetros nominais, inclusive DN métrico contra NPS em polegadas."""
    unit_a, unit_b = norm(unit_a), norm(unit_b)
    value_a, value_b = number(value_a), number(value_b)
    if unit_a == unit_b:
        return math.isclose(value_a, value_b, rel_tol=1e-10, abs_tol=1e-8)
    if {unit_a, unit_b} == {"MM", "POL"}:
        millimetres, inches = ((value_a, value_b) if unit_a == "MM" else (value_b, value_a))
        nominal = next((inch for dn, inch in _DN_TO_INCH.items()
                        if math.isclose(millimetres, dn, rel_tol=0, abs_tol=1e-8)), None)
        return nominal is not None and math.isclose(inches, nominal, rel_tol=0, abs_tol=1e-8)
    try:
        left, right = dimension_to_mm(value_a, unit_a), dimension_to_mm(value_b, unit_b)
    except AnalysisError:
        return False
    tolerance = max(0.6, max(abs(left), abs(right)) * 0.035)
    return abs(left - right) <= tolerance


def _nominal_dimension_convert(value, source_unit, target_unit):
    source_unit, target_unit = norm(source_unit), norm(target_unit)
    value = number(value)
    if source_unit == target_unit:
        return value
    if source_unit == "MM" and target_unit == "POL":
        return next((inch for dn, inch in _DN_TO_INCH.items()
                     if math.isclose(value, dn, rel_tol=0, abs_tol=1e-8)), None)
    if source_unit == "POL" and target_unit == "MM":
        return next((float(dn) for dn, inch in _DN_TO_INCH.items()
                     if math.isclose(value, inch, rel_tol=0, abs_tol=1e-8)), None)
    try:
        return dimension_from_mm(dimension_to_mm(value, source_unit), target_unit)
    except AnalysisError:
        return None


def geometry_mm(dims, units):
    return [dimension_to_mm(value, unit) for value, unit in zip(dims, units)]


def geometry_index(dims, units):
    """Diâmetro equivalente RMS: todas as três bocas influenciam a tendência."""
    values = geometry_mm(dims, units)
    return math.sqrt(sum(value * value for value in values) / len(values))


def format_dimension(value, unit=None, include_unit=False):
    unit = norm(unit)
    if unit == "POL":
        fraction = Fraction(float(value)).limit_denominator(64)
        whole, numerator = divmod(abs(fraction.numerator), fraction.denominator)
        sign = "-" if fraction.numerator < 0 else ""
        if numerator and whole:
            label = f"{sign}{whole} {numerator}/{fraction.denominator}"
        elif numerator:
            label = f"{sign}{numerator}/{fraction.denominator}"
        else:
            label = f"{sign}{whole}"
        return label + ('"' if include_unit else "")
    label = f"{float(value):g}".replace(".", ",")
    if include_unit and unit:
        label += f" {unit.lower()}"
    return label


def format_dimension_input(value, unit):
    """Texto inicial do campo, mantendo frações SINAPI para polegadas."""
    return format_dimension(value, unit, include_unit=(norm(unit) == "POL"))


def format_geometry(dims, units, include_units=True):
    return " × ".join(format_dimension(value, unit, include_unit=include_units)
                      for value, unit in zip(dims, units))


def tee_signature(description):
    geometry = tee_geometry(description)
    return geometry_signature(description, geometry) if geometry else None


def geometry_signature(description, geometry=None):
    geometry = geometry or connection_geometry(description)
    if not geometry:
        return None
    start, end = geometry["span"]
    text = geometry["text"][:start] + " {GEOMETRIA} " + geometry["text"][end:]
    # "TÊ NORMAL" e "TÊ" são equivalentes, mas TÊ DE REDUÇÃO permanece
    # explicitamente separado e pesquisável.
    text = re.sub(r"\bTE\s+NORMAL\b", "TE", text)
    if geometry.get("kind") == "tee":
        text = re.sub(r"\b90\s+GRAUS\b", "", text)
    # Remove somente equivalências nominais em polegadas entre parênteses
    # (ex.: DN 80 (3") ou 65 x 50 mm (2 1/2" x 2")).
    text = re.sub(r"\s*\([^)]*\d[^)]*\sX\s[^)]*\d[^)]*\)", "", text)
    text = re.sub(r'\s*\([^)]*(?:"|POL)[^)]*\)', "", text)
    text = re.sub(r"\s+-?\s*FORNECIMENTO E INSTALACAO\b", " - FORNECIMENTO E INSTALACAO", text)
    text = re.sub(r"\bAF[_ -]?\d{2}/\d{4}\b", "", text)
    text = re.sub(r"\s*[-,.;]+\s*$", "", text)
    text = re.sub(r"\s+,", ",", text)
    return re.sub(r"\s+", " ", text).strip(" ,-.")


def inch_number(text):
    parts = text.split()
    if "/" in parts[-1]:
        numerator, denominator = parts[-1].split("/")
        return (float(parts[0]) if len(parts) > 1 else 0) + float(numerator) / float(denominator)
    return number(text)


def measures(description):
    # Normalizar o expoente antes da transliteração preserva mm².
    raw = str(description).replace("²", "2").replace("³", "3")
    for symbol, replacement in _FRACTIONS.items():
        raw = raw.replace(symbol, replacement)
    text = norm(raw)
    text = re.sub(r'(\d+)[.-](\d+/\d+)(?=\s*(?:"|POL))', r"\1 \2", text)
    # Polegadas entre parênteses após um DN são apenas o segundo nome nominal:
    # DN 80 MM (3") e DN 80 (3") devem usar o mesmo eixo de 80 mm.
    text = re.sub(
        rf'(\b(?:DN|DIAMETRO(?:\s+NOMINAL)?(?:\s+DE)?)\s*=?\s*{_DIM_NUMBER}(?:\s*(?:MM|CM|M))?)'
        r'\s*\([^)]*(?:"|POL)[^)]*\)', r'\1', text, flags=re.I)
    found = []

    def is_range(span):
        # Faixas não são uma única dimensão representativa.
        before = text[max(0, span[0] - 30):span[0]]
        after = text[span[1]:span[1] + 22]
        if re.search(r"(?:ENTRE|ATE|OU|MAIOR QUE|MENOR QUE|SUPERIOR A|INFERIOR A)\s*(?:DN|DE)?\s*$", before):
            return True
        if re.search(r"\d+(?:[.,]\d+)?\s*(?:A|E|ATE)\s*$", before) or re.match(
                r'\s*(?:(?:MM|CM|M|POL(?:EGADAS?)?|\")\b)?\s*(?:OU|A|ATE)\s+(?:(?:DN|DE)\s*)?\d', after):
            return True
        return False

    # Um único DN explícito é o eixo principal, mesmo quando a unidade foi
    # omitida pelo SINAPI (DN 150 = 150 mm). Espessura e peso não viram eixos.
    dn_matches = list(_DN_SEQUENCE.finditer(text))
    if len(dn_matches) == 1:
        dn_match = dn_matches[0]
        parts = list(_DIM_PART.finditer(dn_match["seq"]))
        if len(parts) == 1:
            part = parts[0]
            unit = norm(part["unit"]).replace("POLEGADAS", "POL").replace("POLEGADA", "POL").replace('"', "POL") if part["unit"] else "MM"
            span = (dn_match.start("seq") + part.start("number"),
                    dn_match.start("seq") + part.end("number"))
            value = _parse_dimension_value(part["number"])
            if value > 0 and not is_range(span):
                found.append((span, value, unit))

    if not found:
        for match in MEASURE.finditer(text):
            value = number(match["value"])
            unit = match["unit"].upper()
            if re.fullmatch(r"BTU(?:S)?\s*/?\s*H", unit, re.I):
                unit = "BTU/H"
            if value > 0 and not is_range(match.span("value")):
                found.append((match.span("value"), value, unit))
        for match in INCH.finditer(text):
            value = inch_number(match["value"])
            if value > 0 and not is_range(match.span("value")):
                found.append((match.span("value"), value, "POL"))

    found.sort(key=lambda item: item[0][0])
    result, seen = [], set()
    for span, value, unit in found:
        key = (round(value, 12), unit)
        if key in seen:
            continue
        seen.add(key)
        matches = [s for s, v, u in found if math.isclose(v, value, rel_tol=0, abs_tol=1e-12) and u == unit]
        masked = text
        for start, end in sorted(matches, reverse=True):
            masked = masked[:start] + "{VALOR}" + masked[end:]
        result.append({"value": value, "unit": unit, "template": masked, "start": span[0]})
    return result


def resource_axis_signature(description, axis):
    """Assinatura dimensional sem transformar espessura/peso em nova família.

    Em tubos metálicos, espessura e peso linear mudam naturalmente com o DN.
    Schedule, classe, norma, material e demais especificações permanecem na
    assinatura e, portanto, continuam impedindo agrupamentos indevidos.
    """
    text = axis["template"]
    source = norm(description)
    if re.search(r"\bTUBOS?\b", source) and re.search(r"\b(?:ACO|FERRO)\b", source):
        text = re.sub(
            r"\s*,?\s*\b(?:E|ESPESSURA)\s*=?\s*\*?\s*\d+(?:[.,]\d+)?\s*\*?\s*MM\b",
            "", text)
        text = re.sub(
            r"\s*,?\s*(?:PESO\s*)?\(?\s*\*?\s*\d+(?:[.,]\d+)?\s*\*?\s*KG/M\s*\)?",
            "", text)
        text = re.sub(r"\s+,", ",", text)
        text = re.sub(r",\s*,", ",", text)
    return re.sub(r"\s+", " ", text).strip(" ,")


def dimension_range(description):
    """Lê faixas oficiais como DN 1000 mm OU DN 1100 mm."""
    text = _normalized_dimensions_text(description)
    pattern = re.compile(
        rf"\bDN\s*=?\s*(?P<first>{_DIM_NUMBER})\s*(?P<unit1>{_DIM_UNIT})?"
        rf"\s*(?:\([^)]*\))?\s+OU\s+(?:DN\s*=?\s*)?"
        rf"(?P<second>{_DIM_NUMBER})\s*(?P<unit2>{_DIM_UNIT})?\s*(?:\([^)]*\))?",
        re.I)
    matches = list(pattern.finditer(text))
    if len(matches) != 1:
        return None
    match = matches[0]
    first, second = _parse_dimension_value(match["first"]), _parse_dimension_value(match["second"])
    unit1 = norm(match["unit1"]).replace("POLEGADAS", "POL").replace("POLEGADA", "POL").replace('"', "POL") if match["unit1"] else ""
    unit2 = norm(match["unit2"]).replace("POLEGADAS", "POL").replace("POLEGADA", "POL").replace('"', "POL") if match["unit2"] else ""
    inferred = unit1 or unit2 or "MM"
    unit1, unit2 = unit1 or inferred, unit2 or inferred
    if first <= 0 or second <= 0 or unit1 != unit2 or math.isclose(first, second):
        return None
    low, high = sorted((first, second))
    signature = text[:match.start()] + "DN {FAIXA}" + text[match.end():]
    signature = re.sub(r"\s+", " ", signature).strip()
    return {"low": low, "high": high, "unit": unit1, "x": (low + high) / 2,
            "span": match.span(), "text": text, "signature": signature}


def format_reference_dimensions(ref):
    if ref.get("range"):
        low, high = ref["range"]
        unit = ref["units"][0]
        return (f"{format_dimension(low, unit, include_unit=False)}–"
                f"{format_dimension(high, unit, include_unit=True)}")
    return format_geometry(ref["dims"], ref["units"])


def discipline(description, group=""):
    text = norm(description + " " + group)
    if is_cable_tray(description) or "ELETROCALHAS" in norm(group):
        return "Elétrica"
    if any(x in text for x in ("INCENDIO", "HIDRANTE", "SPRINKLER", "EXTINTOR")):
        return "PCI / incêndio"
    if any(x in text for x in ("AR CONDICIONADO", "CLIMATIZ", "REFRIGER", "DUTO DE AR")):
        return "Climatização"
    if any(x in text for x in ("GAS COMBUSTIVEL", "GAS MEDICINAL")):
        return "Gases / tubulações"
    if any(x in text for x in ("ESGOTO", "AGUA PLUVIAL", "DRENAGEM")):
        return "Esgoto e drenagem"
    return "Água / tubulações"


def build_families(database):
    grouped = defaultdict(list)
    for code, composition in database["compositions"].items():
        if (not composition["components"] or "EM ESTUDO" in norm(composition["status"])
                or not is_automatic_scope(composition["description"])):
            continue
        if is_cable_tray(composition["description"]):
            geometry = cable_tray_geometry(composition["description"])
            if not geometry:
                continue
            key = ("cable_tray", norm(composition["group"]), norm(composition["unit"]),
                   tuple(geometry["units"]), geometry["kind"],
                   geometry["profile_offset_mm"], cable_tray_signature(composition["description"], geometry))
            grouped[key].append({"code": code, "dims": geometry["dims"], "units": geometry["units"],
                "x": (geometry_index(geometry["dims"], geometry["units"])
                      if len(geometry["dims"]) > 1 else geometry["dims"][0]),
                "reported_dimensions": len(geometry["dims"]), "inferred": False,
                "geometry_kind": geometry["kind"]})
            continue
        geometry = connection_geometry(composition["description"])
        if geometry:
            signature = family_signature_text(geometry_signature(composition["description"], geometry))
            key = ("geometry", norm(composition["group"]), norm(composition["unit"]),
                   len(geometry["dims"]), tuple(geometry["units"]), signature)
            grouped[key].append({"code": code, "dims": geometry["dims"], "units": geometry["units"],
                "x": geometry_index(geometry["dims"], geometry["units"]),
                "reported_dimensions": geometry["reported_dimensions"], "inferred": geometry["inferred"],
                "geometry_kind": geometry["kind"]})
            continue
        interval = dimension_range(composition["description"])
        if interval:
            key = ("range", norm(composition["group"]), norm(composition["unit"]),
                   interval["unit"], family_signature_text(interval["signature"]))
            grouped[key].append({"code": code, "x": interval["x"], "dims": [interval["x"]],
                "units": [interval["unit"]], "range": [interval["low"], interval["high"]],
                "reported_dimensions": 2, "inferred": False})
            continue
        axes = measures(composition["description"])
        # Uma CPU de tubo/conexão sem geometria X possui um único DN principal.
        # Se a descrição também informar espessura/peso, o primeiro eixo textual
        # é o nominal e os demais não geram famílias paralelas.
        axes = axes[:1]
        for axis in axes:
            key = ("single", norm(composition["group"]), norm(composition["unit"]), axis["unit"],
                   family_signature_text(axis["template"]))
            grouped[key].append({"code": code, "x": axis["value"], "dims": [axis["value"]],
                                 "units": [axis["unit"]], "reported_dimensions": 1, "inferred": False})
    families = []
    for key, refs in grouped.items():
        tray_details = {}
        if key[0] == "cable_tray":
            _, group, service_unit, stored_units, geometry_kind, profile_offset_mm, title = key
            axis_units, dimension_count = list(stored_units), len(stored_units)
            unit = axis_units[0] if len(set(axis_units)) == 1 else "MISTA"
            axis_labels = (["Largura atendida", "Comprimento do perfil"] if geometry_kind == "tray_support"
                           else ["Largura maior", "Largura menor", "Altura"] if geometry_kind == "tray_reduction"
                           else ["Altura"] if geometry_kind == "tray_height_accessory"
                           else ["Largura", "Altura"])
            for i, label in enumerate(axis_labels):
                title = title.replace(f"{{D{i + 1}}}", "{" + label.upper() + "} " + axis_units[i])
            curve_label = (f"Índice RMS das {dimension_count} dimensões (mm)" if dimension_count > 1
                           else f"{axis_labels[0]} ({unit.lower()})")
            tray_details = {"cable_tray": True, "geometry_kind": geometry_kind, "axis_labels": axis_labels,
                            "height_axis": (None if geometry_kind == "tray_support" else dimension_count - 1),
                            "profile_offset_mm": profile_offset_mm}
        elif key[0] == "geometry":
            _, group, service_unit, dimension_count, stored_units, signature = key
            axis_units = list(stored_units)
            placeholders = " × ".join(f"{{D{i + 1}}}" for i in range(dimension_count))
            title = signature.replace("{GEOMETRIA}", placeholders)
            unit = axis_units[0] if len(set(axis_units)) == 1 else "MISTA"
            curve_label = f"Diâmetro equivalente das {dimension_count} dimensões (mm)"
        elif key[0] == "range":
            _, group, service_unit, unit, title = key
            axis_units, dimension_count = [unit], 1
            curve_label = f"Ponto médio da faixa DN ({unit})"
        else:
            _, group, service_unit, unit, title = key
            axis_units, dimension_count = [unit], 1
            curve_label = f"Dimensão ({unit})"
        candidate = {"id": hashlib.sha1(repr(key).encode()).hexdigest()[:14], "title": title,
            "group": group, "unit": unit, "axis_units": axis_units, "dimension_count": dimension_count,
            "curve_label": curve_label, "service_unit": service_unit, "manual": False,
            "tee": bool(refs and refs[0].get("geometry_kind") == "tee"),
            "range_based": key[0] == "range",
            "discipline": discipline(title, group),
            "refs": sorted(refs, key=lambda r: (r["x"], tuple(r["dims"]), r["code"]))}
        candidate.update(tray_details)
        # Não exibir famílias que só parecem dimensionais, mas não possuem ao
        # menos dois pontos com o mesmo conjunto de recursos da CPU.
        try:
            maps = aligned_resources(database, candidate["refs"], candidate)
        except AnalysisError:
            continue
        roster_xs = defaultdict(set)
        for ref, mapping in zip(candidate["refs"], maps):
            roster_xs[frozenset(mapping)].add(round(ref["x"], 10))
        candidate["has_trend"] = any(len(xs) >= 2 for xs in roster_xs.values())
        candidate["single_reference"] = len({ref["code"] for ref in refs}) == 1
        metal_family = bool(re.search(r"\b(?:ACO|FERRO)\b", norm(title)))
        if not candidate["has_trend"] and not metal_family and not candidate["single_reference"]:
            continue
        families.append(candidate)
    return sorted(families, key=lambda f: (f["discipline"], f["title"]))


def manual_family(database, pairs, unit, title="Família selecionada pelo usuário"):
    refs = []
    compositions = database["compositions"]
    dimension_count = None
    raw_units = [norm(value) for value in (unit if isinstance(unit, (list, tuple)) else str(unit).split(",")) if str(value).strip()]
    for pair in pairs:
        if len(pair) < 2:
            raise AnalysisError("Informe código e dimensão em cada linha.")
        code, dimensions = pair[0], pair[1:]
        code = str(code).strip()
        if code not in compositions:
            raise AnalysisError(f"Código {code} não encontrado na base ativa.")
        dims = [dimension_number(value) for value in dimensions]
        dimension_count = dimension_count or len(dims)
        if len(dims) != dimension_count or dimension_count not in (1, 2, 3):
            raise AnalysisError("Use uma dimensão para tubos/conexões comuns, duas para conexões de redução ou três para tês.")
        if any(value <= 0 for value in dims):
            raise AnalysisError("As dimensões devem ser maiores que zero.")
        units = raw_units if len(raw_units) == dimension_count else raw_units * dimension_count
        if len(units) != dimension_count:
            raise AnalysisError("Informe uma unidade ou uma unidade para cada dimensão.")
        x = geometry_index(dims, units) if dimension_count > 1 else dims[0]
        refs.append({"code": code, "x": x, "dims": dims, "units": units,
                     "reported_dimensions": dimension_count, "inferred": False})
    if len(refs) < 2:
        raise AnalysisError("Informe ao menos 2 composições comparáveis.")
    units = {norm(compositions[r["code"]]["unit"]) for r in refs}
    if len(units) != 1:
        raise AnalysisError("As composições precisam usar a mesma unidade de serviço.")
    axis_units = refs[0]["units"]
    display_unit = axis_units[0] if len(set(axis_units)) == 1 else "MISTA"
    return {"id": "manual", "title": title, "group": "Seleção manual", "unit": display_unit,
            "axis_units": axis_units, "dimension_count": dimension_count,
            "curve_label": (f"Diâmetro equivalente das {dimension_count} dimensões (mm)"
                            if dimension_count > 1 else f"Dimensão ({display_unit})"),
            "service_unit": units.pop(), "discipline": "Seleção manual", "manual": True,
            "has_trend": True, "range_based": False,
            "tee": dimension_count == 3,
            "refs": sorted(refs, key=lambda r: (r["x"], tuple(r["dims"])))}


def pchip_slopes(xs, ys):
    if len(xs) < 2 or len(xs) != len(ys) or any(b <= a for a, b in zip(xs, xs[1:])):
        raise AnalysisError("A curva exige dimensões distintas em ordem crescente.")
    h = [b - a for a, b in zip(xs, xs[1:])]
    d = [(b - a) / step for a, b, step in zip(ys, ys[1:], h)]
    if len(xs) == 2:
        return [d[0], d[0]]
    slopes = [0.0] * len(xs)
    for k in range(1, len(xs) - 1):
        if d[k - 1] * d[k] > 0:
            w1, w2 = 2 * h[k] + h[k - 1], h[k] + 2 * h[k - 1]
            slopes[k] = (w1 + w2) / (w1 / d[k - 1] + w2 / d[k])
    def end(h0, h1, d0, d1):
        value = ((2 * h0 + h1) * d0 - h0 * d1) / (h0 + h1)
        if value * d0 <= 0:
            return 0.0
        if d0 * d1 < 0 and abs(value) > abs(3 * d0):
            return 3 * d0
        return value
    slopes[0] = end(h[0], h[1], d[0], d[1])
    slopes[-1] = end(h[-1], h[-2], d[-1], d[-2])
    return slopes


def pchip(xs, ys, target):
    slopes = pchip_slopes(xs, ys)
    if target < xs[0] or target > xs[-1]:
        endpoint = 0 if target < xs[0] else -1
        return ys[endpoint] + slopes[endpoint] * (target - xs[endpoint])
    k = min(len(xs) - 2, max(0, bisect.bisect_right(xs, target) - 1))
    h = xs[k + 1] - xs[k]
    t = (target - xs[k]) / h
    return ((2 * t**3 - 3 * t**2 + 1) * ys[k] + (t**3 - 2 * t**2 + t) * h * slopes[k]
            + (-2 * t**3 + 3 * t**2) * ys[k + 1] + (t**3 - t**2) * h * slopes[k + 1])


def validation(xs, ys):
    # Retira cada ponto INTERNO; endpoints não têm observações dos dois lados.
    if len(xs) < 4:
        label = ("2 pontos: tendência linear com maior possibilidade de erro" if len(xs) == 2
                 else "3 pontos: sem validação interna suficiente")
        return {"n": 0, "nmae_pct": None, "label": label}
    errors = []
    for k in range(1, len(xs) - 1):
        value = pchip(xs[:k] + xs[k + 1:], ys[:k] + ys[k + 1:], xs[k])
        errors.append(abs(value - ys[k]))
    scale = sum(abs(y) for y in ys[1:-1]) / len(errors)
    score = (sum(errors) / len(errors) / scale * 100) if scale else (0 if not any(errors) else None)
    label = "Variação irregular: revisar" if score is None or score > 20 else "Validação interna disponível"
    return {"n": len(errors), "nmae_pct": score, "label": label}


def resource_kind(component):
    description, unit = norm(component["description"]), norm(component["unit"])
    if unit in ("CHP", "CHI"):
        return "Equipamento"
    if unit == "H" and ("ENCARGOS COMPLEMENTARES" in description or any(w in description for w in ("AJUDANTE", "ENCANADOR", "ELETRICISTA", "SERVENTE", "PEDREIRO"))):
        return "Mão de obra"
    return "Composição auxiliar" if "COMPOS" in component["kind"] else "Material / insumo"


def aligned_resources(database, refs, family):
    codesets = [{(c["code"], norm(c["unit"]), c["kind"]) for c in database["compositions"][r["code"]]["components"]} for r in refs]
    common = set.intersection(*codesets) if codesets else set()
    maps = []
    for ref in refs:
        mapping = {}
        for component in database["compositions"][ref["code"]]["components"]:
            identity = (component["code"], norm(component["unit"]), component["kind"])
            if identity in common:
                key = ("codigo",) + identity
            elif family.get("cable_tray") and is_cable_tray(component["description"]):
                geometry = cable_tray_geometry(component["description"])
                if geometry:
                    key = ("eletrocalha", cable_tray_signature(component["description"], geometry, align_gauge=True),
                           geometry["kind"], geometry["profile_offset_mm"],
                           norm(component["unit"]), component["kind"])
                else:
                    key = ("codigo",) + identity
            elif family["dimension_count"] > 1 and connection_geometry(component["description"]):
                component_geometry = connection_geometry(component["description"])
                if len(component_geometry["dims"]) == family["dimension_count"]:
                    key = ("geometria", geometry_signature(component["description"], component_geometry),
                           norm(component["unit"]), component["kind"])
                else:
                    key = ("codigo",) + identity
            else:
                all_axes = measures(component["description"])
                matching = [axis for axis in all_axes
                            if nominal_dimensions_equal(axis["value"], axis["unit"],
                                                        ref["dims"][0], family["axis_units"][0])]
                # O serviço pode usar polegadas e o insumo correspondente, mm nominais.
                # Um único eixo inequívoco na descrição continua sendo dimensional.
                axes = matching if len(matching) == 1 else all_axes if len(all_axes) == 1 else []
                key = ("dimensao", resource_axis_signature(component["description"], axes[0]),
                       norm(component["unit"]), component["kind"]) if len(axes) == 1 else ("codigo",) + identity
            if key in mapping:
                raise AnalysisError(f"Componente repetido/ambíguo na composição {ref['code']}: {component['code']}. Não somei linhas automaticamente.")
            mapping[key] = component
        maps.append(mapping)
    return maps


def _target_geometry(family, target):
    if family["dimension_count"] == 1:
        value = dimension_number(target[0] if isinstance(target, (list, tuple)) else target)
        if value <= 0:
            raise AnalysisError("Informe uma dimensão positiva.")
        return [value], list(family["axis_units"]), value
    values = list(target) if isinstance(target, (list, tuple)) else re.split(r"\s*[xX;|]\s*", str(target).strip())
    count = family["dimension_count"]
    if len(values) != count:
        raise AnalysisError(f"Informe as {count} dimensões do serviço.")
    dims = [dimension_number(value) for value in values]
    if any(value <= 0 for value in dims):
        raise AnalysisError("Todas as dimensões devem ser positivas.")
    units = list(family["axis_units"])
    if family.get("geometry_kind") == "tray_support":
        offset = dimension_to_mm(dims[1], units[1]) - dimension_to_mm(dims[0], units[0])
        if not math.isclose(offset, family["profile_offset_mm"], abs_tol=1e-8):
            raise AnalysisError("Nos suportes desta família, o comprimento do perfil deve ser a largura atendida "
                                f"mais {family['profile_offset_mm']:g} mm, conforme as referências.")
    if family.get("geometry_kind") == "tray_reduction" and dims[0] < dims[1]:
        raise AnalysisError("Informe a largura maior primeiro e a largura menor depois.")
    return dims, units, geometry_index(dims, units)


def _same_target(ref, dims, units):
    if len(ref["dims"]) != len(dims):
        return False
    if ref.get("range") and len(dims) == 1:
        try:
            target = dimension_from_mm(dimension_to_mm(dims[0], units[0]), ref["units"][0])
        except AnalysisError:
            return False
        return ref["range"][0] - 1e-8 <= target <= ref["range"][1] + 1e-8
    if len(dims) == 1:
        return math.isclose(ref["dims"][0], dims[0], rel_tol=1e-12, abs_tol=1e-10)
    left, right = geometry_mm(ref["dims"], ref["units"]), geometry_mm(dims, units)
    return all(math.isclose(a, b, rel_tol=1e-10, abs_tol=1e-8) for a, b in zip(left, right))


def _curve_points(points):
    """Agrupa índices equivalentes sem inventar uma ordem entre permutações de bocas."""
    groups = defaultdict(list)
    for point in points:
        groups[round(point["x"], 10)].append(point)
    xs, ys = [], []
    for x in sorted(groups):
        values = [point["coefficient"] for point in groups[x]]
        xs.append(float(x))
        ys.append(None if any(value is None for value in values) else sum(values) / len(values))
    return xs, ys


def _supply_axis_signature(description, family, source_dims=None):
    if family.get("cable_tray"):
        geometry = cable_tray_geometry(description)
        return cable_tray_signature(description, geometry) if geometry else None
    if family["dimension_count"] > 1:
        geometry = connection_geometry(description)
        if geometry and len(geometry["dims"]) == family["dimension_count"]:
            return geometry_signature(description, geometry)
        return None
    axes = measures(description)
    if source_dims:
        selected = [axis for axis in axes
                    if nominal_dimensions_equal(axis["value"], axis["unit"],
                                                source_dims[0], family["axis_units"][0])]
        if len(selected) == 1:
            return resource_axis_signature(description, selected[0])
    return resource_axis_signature(description, axes[0]) if len(axes) == 1 else None


def _supply_matches(database, base_component, exact_component, family, source_ref, target_dims, target_units):
    supplies = database.get("supplies", {})
    if exact_component and exact_component.get("code") in supplies:
        match = supplies[exact_component["code"]]
        return {"checked": True, "matches": [match], "status": _supply_price_status([match])}
    if not base_component:
        return {"checked": False, "matches": [], "status": "Não se aplica a este componente"}
    if base_component.get("code") not in supplies:
        if "INSUMO" in norm(base_component.get("kind")) and is_automatic_scope(base_component.get("description")):
            return {"checked": True, "matches": [],
                    "status": "Insumo do serviço não consta com preço na aba ISD"}
        return {"checked": False, "matches": [], "status": "Não se aplica a este componente"}
    source_supply = supplies[base_component["code"]]
    signature = _supply_axis_signature(source_supply["description"], family, source_ref["dims"])
    matches = []
    if signature:
        for supply in supplies.values():
            candidate_dims = target_dims if family["dimension_count"] == 1 else source_ref["dims"]
            if _supply_axis_signature(supply["description"], family, candidate_dims) != signature:
                continue
            if family.get("cable_tray"):
                geometry = cable_tray_geometry(supply["description"])
                wanted_dims, wanted_units = target_dims, target_units
                if geometry and geometry["kind"] == "tray_height_accessory" and family.get("height_axis") is not None:
                    axis = family["height_axis"]
                    wanted_dims, wanted_units = [target_dims[axis]], [target_units[axis]]
                if (geometry and len(geometry["dims"]) == len(wanted_dims)
                        and all(math.isclose(a, b, abs_tol=1e-8) for a, b in zip(
                            geometry_mm(geometry["dims"], geometry["units"]), geometry_mm(wanted_dims, wanted_units)))):
                    matches.append(supply)
            elif family["dimension_count"] > 1:
                geometry = connection_geometry(supply["description"])
                if geometry and _same_target({"dims": geometry["dims"], "units": geometry["units"]}, target_dims, target_units):
                    matches.append(supply)
            else:
                def equivalent_axis(axis):
                    return nominal_dimensions_equal(axis["value"], axis["unit"],
                                                    target_dims[0], family["axis_units"][0])
                if any(equivalent_axis(axis) for axis in measures(supply["description"])):
                    matches.append(supply)
    matches.sort(key=lambda item: item["code"])
    return {"checked": True, "matches": matches, "status": _supply_price_status(matches)}


def _supply_price_status(matches):
    if not matches:
        return "Insumo SINAPI da dimensão desejada não encontrado"
    states = sorted({state for item in matches for state in ("PE", "SP") if item.get(state) is not None})
    if not states:
        return "Insumo encontrado, sem preço em PE ou SP"
    return "Insumo com preço em " + " e ".join(states)


def _estimated_target_description(base, family, target_dims, target_units):
    if not base:
        return ""
    if family.get("cable_tray") and is_cable_tray(base["description"]):
        geometry = cable_tray_geometry(base["description"])
        if not geometry or (geometry["kind"] == "tray_support" and family["geometry_kind"] != "tray_support"):
            # O suporte 400/800 atende várias larguras de eletrocalha. A largura
            # da eletrocalha não pode substituir a especificação do suporte.
            return base["description"] + " [SELECIONAR E CONFERIR O SUPORTE DO ALVO]"
        wanted_dims, wanted_units = target_dims, target_units
        if geometry["kind"] == "tray_height_accessory" and family.get("height_axis") is not None:
            axis = family["height_axis"]
            wanted_dims, wanted_units = [target_dims[axis]], [target_units[axis]]
        if len(geometry["dims"]) != len(wanted_dims):
            return base["description"] + " [ESPECIFICAÇÃO A CONFERIR]"
        text = geometry["text"]
        for index in reversed(range(len(geometry["spans"]))):
            start, end = geometry["spans"][index]
            value = dimension_from_mm(dimension_to_mm(wanted_dims[index], wanted_units[index]), geometry["units"][index])
            text = text[:start] + format_dimension(value, geometry["units"][index], include_unit=True).upper() + text[end:]
        text, gauge_count = re.subn(r"\bESPESSURA\s*#\s*\d+|#\s*\d+", "ESPESSURA A DEFINIR", text)
        return text + (" [ESPESSURA A DEFINIR; ESPECIFICAÇÃO A CONFERIR]" if gauge_count
                       else " [ESPECIFICAÇÃO A CONFERIR]")
    if family["dimension_count"] > 1:
        geometry = connection_geometry(base["description"])
        if geometry:
            start, end = geometry["span"]
            target = ("DN " if geometry.get("has_dn") else "") + format_geometry(target_dims, target_units)
            return geometry["text"][:start] + target + geometry["text"][end:] + " [ESPECIFICAÇÃO A CONFERIR]"
    axes = measures(base["description"])
    if axes:
        axis = axes[0]
        target_value = _nominal_dimension_convert(target_dims[0], target_units[0], axis["unit"])
        if target_value is not None:
            label = format_dimension(target_value, axis["unit"])
            return axis["template"].replace("{VALOR}", label) + " [ESPECIFICAÇÃO A CONFERIR]"
        # Não transforme um DN métrico sem NPS oficial em uma fração arbitrária.
        # Em tubos metálicos, remova também espessura/peso herdados do tubo-base.
        template = resource_axis_signature(base["description"], axis)
        label = format_dimension(target_dims[0], target_units[0], include_unit=True)
        template, replacements = re.subn(
            r'\{VALOR\}\s*(?:"|POL(?:EGADAS?)?\b)', label, template, count=1)
        if not replacements:
            template = template.replace("{VALOR}", label)
        template = re.sub(r"(?<=\w)\(", " (", template)
        suffix = (" [ESPESSURA E PESO A DEFINIR; ESPECIFICAÇÃO A CONFERIR]"
                  if re.search(r"\bTUBOS?\b", norm(base["description"]))
                  and re.search(r"\b(?:ACO|FERRO)\b", norm(base["description"]))
                  else " [ESPECIFICAÇÃO A CONFERIR]")
        return template + suffix
    return base["description"]


def _tray_height_resource_changes(component, family, target_dims, target_units):
    axis = family.get("height_axis")
    if not component or not family.get("cable_tray") or axis is None:
        return False
    geometry = cable_tray_geometry(component["description"])
    return bool(geometry and geometry["height_mm"] is not None
                and not math.isclose(geometry["height_mm"],
                                     dimension_to_mm(target_dims[axis], target_units[axis]), abs_tol=1e-8))


def calculate(database, family, source_code, target, allow_extrapolation=False):
    target_dims, target_units, target_x = _target_geometry(family, target)
    refs = family["refs"]
    source_index = next((i for i, r in enumerate(refs) if r["code"] == source_code), None)
    if source_index is None:
        raise AnalysisError("Selecione uma composição de origem desta família.")
    maps = aligned_resources(database, refs, family)
    exact_index = next((i for i, ref in enumerate(refs) if _same_target(ref, target_dims, target_units)), None)
    roster = set(maps[source_index])
    compatible_indices = [i for i, mapping in enumerate(maps) if set(mapping) == roster]
    segment_refs = [refs[i] for i in compatible_indices]
    segment_maps = [maps[i] for i in compatible_indices]
    unique_xs = sorted({round(ref["x"], 10) for ref in segment_refs})
    if exact_index is None and len(unique_xs) < 2:
        raise AnalysisError("Há menos de 2 geometrias compatíveis com a composição de origem.")
    index_extrapolated = bool(unique_xs and (target_x < unique_xs[0] or target_x > unique_xs[-1]))
    outside_axes = []
    if family.get("cable_tray") and exact_index is None:
        # O índice RMS pode estar dentro da curva mesmo que a altura esteja
        # fora das alturas observadas. Isso continua sendo extrapolação da
        # geometria e exige a opção explícita de extrapolação.
        for axis, wanted in enumerate(geometry_mm(target_dims, target_units)):
            observed = [geometry_mm(ref["dims"], ref["units"])[axis] for ref in segment_refs]
            if observed and (wanted < min(observed) - 1e-8 or wanted > max(observed) + 1e-8):
                outside_axes.append(axis)
    extrapolated = index_extrapolated or bool(outside_axes)
    warnings = []
    if family.get("cable_tray"):
        height_axis = family.get("height_axis")
        if height_axis is not None:
            observed_heights = sorted({geometry_mm(ref["dims"], ref["units"])[height_axis] for ref in segment_refs})
            heights_text = ", ".join(f"{height:g}" for height in observed_heights)
            warnings.append(f"Alturas observadas nas referências utilizadas: {heights_text} mm. "
                            "O índice RMS considera largura(s) e altura em conjunto; "
                            "não identifica uma tendência independente para a altura.")
            if height_axis in outside_axes:
                target_height = geometry_mm(target_dims, target_units)[height_axis]
                warnings.append(f"ALTURA SEM REFERÊNCIA OBSERVADA: alvo de {target_height:g} mm; "
                                f"referências com altura(s) de {heights_text} mm. "
                                "A adaptação usa uma hipótese dimensional pelo índice RMS, "
                                "sem coeficientes observados para essa altura; revisão técnica obrigatória.")
        if family.get("geometry_kind") == "tray_support":
            warnings.append("Suportes: largura atendida e comprimento do perfil variam juntos nas referências; "
                            f"mantém-se a diferença de {family['profile_offset_mm']:g} mm. "
                            "A tendência não verifica a capacidade estrutural do suporte.")
        if exact_index is None:
            warnings.append("Os materiais de eletrocalha são alinhados pela função, com as trocas de espessura "
                            "observadas nas referências. A espessura da chapa e os suportes do alvo devem ser conferidos; "
                            "não são dimensionados pela curva de coeficientes.")
    if family["manual"]:
        warnings.append("Comparabilidade técnica da família declarada pelo usuário; confira material, junta, aplicação, classe e método executivo.")
    if len(segment_refs) < len(refs):
        warnings.append("A lista de componentes muda na família. A curva usa somente as referências com o mesmo conjunto de componentes da origem.")
    if family.get("tee") and any(ref.get("inferred") for ref in refs):
        warnings.append("Tês com uma dimensão foram interpretados como D × D × D; com duas, como D1 × D1 × D2. A descrição SINAPI original permanece nas referências.")
    if family.get("range_based"):
        warnings.append("Composições SINAPI definidas por faixa DN usam o ponto médio somente como eixo da tendência. Um alvo dentro de uma faixa existente é consultado diretamente, sem estimativa.")
    if exact_index is not None:
        mode, method = "Referência SINAPI existente", "Consulta direta; sem estimativa"
        target_map = maps[exact_index]
        keys = list(target_map) + [k for k in maps[source_index] if k not in target_map]
    else:
        if len(unique_xs) == 2:
            warnings.append("ATENÇÃO: há somente 2 pontos compatíveis. O cálculo é linear entre eles e possui maior possibilidade de erro; faça revisão técnica obrigatória.")
        if extrapolated:
            if not allow_extrapolation:
                axis_text = ", ".join(family["axis_labels"][axis].lower() for axis in outside_axes)
                detail = f" ({axis_text})" if axis_text else ""
                raise AnalysisError(f"O alvo está fora do intervalo observado{detail}. Marque a opção de extrapolação para continuar.")
            mode = "Estimativa própria — extrapolação"
            if index_extrapolated:
                method = ("Extrapolação linear por componente (2 pontos)" if len(unique_xs) == 2
                          else "Extrapolação pela tangente da extremidade PCHIP")
            else:
                method = ("Extrapolação dimensional via interpolação linear por componente (2 pontos)"
                          if len(unique_xs) == 2 else "Extrapolação dimensional via PCHIP")
            warnings.append("EXTRAPOLAÇÃO SEM LIMITE PERCENTUAL: não há observações envolvendo o alvo. Revise obrigatoriamente todos os coeficientes antes de usar.")
        else:
            mode, method = "Estimativa própria — interpolação", "PCHIP por componente" if len(unique_xs) >= 3 else "Interpolação linear por componente (2 pontos)"
        if family["dimension_count"] > 1:
            index_label = "índice RMS" if family.get("cable_tray") else "diâmetro equivalente RMS"
            method += f"; eixo = {index_label} das {family['dimension_count']} dimensões"
        keys = list(maps[source_index])
        target_map = None
    result_rows = []
    for key in keys:
        base = maps[source_index].get(key)
        points = []
        for ref, mapping in zip(segment_refs, segment_maps):
            component = mapping.get(key)
            if component:
                points.append({"x": ref["x"], "coefficient": component["coefficient"],
                    "dimensions": ref["dims"], "dimension_text": format_reference_dimensions(ref),
                    "composition": ref["code"], "resource": component["code"], "row": component["row"]})
        reference = (target_map.get(key) if target_map is not None else base)
        current = reference or base
        xs, ys = _curve_points(points)
        check = {"n": 0, "nmae_pct": None, "label": "Consulta direta"}
        if target_map is not None:
            value = reference["coefficient"] if reference else None
            status = "Oficial na referência" if reference and value is not None else ("Não existe na composição de destino" if not reference else "Coeficiente ausente na referência")
            target_resource = reference["code"] if reference else ""
            target_description = reference["description"] if reference else ""
        elif any(y is None for y in ys):
            value, status, target_resource, target_description = None, "Sem dados completos; não estimado", "", ""
        else:
            check = validation(xs, ys)
            value = pchip(xs, ys, target_x)
            status = "Estimado; revisar especificação"
            if check["nmae_pct"] is not None and check["nmae_pct"] > 20:
                status = "Estimado; variação irregular — revisar"
            if not math.isfinite(value) or value < 0:
                value, status = None, "Tendência inválida/negativa; não estimado"
            target_resource = base["code"] if key[0] == "codigo" else "DEFINIR INSUMO DO ALVO"
            target_description = base["description"] if key[0] == "codigo" else _estimated_target_description(base, family, target_dims, target_units)
        height_resource_changes = (target_map is None
                                   and _tray_height_resource_changes(base, family, target_dims, target_units))
        if height_resource_changes and key[0] == "codigo":
            target_resource = "DEFINIR INSUMO DO ALVO"
            target_description = _estimated_target_description(base, family, target_dims, target_units)
        exact_component = target_map.get(key) if target_map is not None else None
        if target_map is None and family.get("cable_tray") and key[0] == "codigo" and not height_resource_changes:
            exact_component = base
        supply = _supply_matches(database, base, exact_component,
                                 family, refs[source_index], target_dims, target_units)
        if supply["matches"] and target_map is None and (key[0] != "codigo" or height_resource_changes):
            target_resource = ", ".join(item["code"] for item in supply["matches"])
            target_description = supply["matches"][0]["description"]
        original = base["coefficient"] if base else None
        change = (value / original - 1) * 100 if value is not None and original not in (None, 0) else None
        result_rows.append({"kind": resource_kind(current), "source_resource": base["code"] if base else "",
            "description": base["description"] if base else current["description"],
            "target_description": target_description, "target_resource": target_resource,
            "unit": current["unit"], "source_coefficient": original, "coefficient": value,
            "change_pct": change, "status": status, "validation": check, "points": points,
            "supply_checked": supply["checked"], "supply_status": supply["status"],
            "supply_matches": supply["matches"]})
    if exact_index is None and any(r["coefficient"] is None for r in result_rows):
        warnings.append("Resultado parcial: existem componentes sem estimativa. A CPU não está completa para uso no orçamento.")
    if exact_index is None:
        warnings.append("Coeficientes estimados não são oficiais SINAPI e não criam código oficial. Os preços PE/SP exibidos são apenas dos insumos encontrados na referência.")
    source_text = format_reference_dimensions(refs[source_index])
    target_text = format_geometry(target_dims, target_units)
    used_refs = [refs[exact_index]] if exact_index is not None else segment_refs
    height_note = ""
    if family.get("height_axis") in outside_axes:
        height_note = "; altura sem referência observada, estimada por hipótese dimensional RMS; revisão técnica obrigatória"
    justification = build_justification(
        'SINAPI', database['meta']['period'],
        [f"{ref['code']} ({format_reference_dimensions(ref)})" for ref in used_refs],
        f'{method}, com origem na referência {source_code}, de {source_text} para {target_text}{height_note}')
    return {"version": VERSION, "created_at": datetime.now().isoformat(timespec="seconds"), "meta": database["meta"],
        "family": family, "source_code": source_code, "source_description": database["compositions"][source_code]["description"],
        "source_x": refs[source_index]["x"], "source_dims": refs[source_index]["dims"], "source_dimension_text": source_text,
        "target": target_x, "target_dims": target_dims, "target_dimension_text": target_text,
        "target_code": refs[exact_index]["code"] if exact_index is not None else None,
        "mode": mode, "method": method, "justification": justification, "warnings": warnings,
        "range": [unique_xs[0], unique_xs[-1]] if unique_xs else [target_x, target_x],
        "references": [{**r, "description": database["compositions"][r["code"]]["description"],
                        "dimension_text": format_reference_dimensions(r),
                        "row": database["compositions"][r["code"]]["row"], "used": i in compatible_indices} for i, r in enumerate(refs)],
        "rows": result_rows}


def export_report(result, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    summary = book.active
    summary.title = "Leia primeiro"
    metadata = [("PROCESSO 3 — COEFICIENTES SINAPI", VERSION), ("Resultado", result["mode"]),
        ("Banco", result['meta'].get('bank') or 'SINAPI'),
        ("Competência", result["meta"]["period"]), ("Método", result["method"]),
        ("Origem", result["source_code"]), ("Descrição de origem", result["source_description"]),
        ("Dimensão original", result["source_dimension_text"]), ("Dimensão alvo", result["target_dimension_text"]),
        ("Eixo matemático", result["family"]["curve_label"]), ("Unidade do serviço", result["family"]["service_unit"]),
        ("Código oficial da composição no alvo", result["target_code"] or "Não atribuído: estimativa própria, não é novo código SINAPI"),
        ("Código da composição de origem", result["source_code"]),
        ("Justificativa curta", result["justification"]),
        ("Arquivo de referência", result["meta"]["filename"]), ("Aba fonte", result["meta"]["sheet"]),
        ("Aba de preços", result["meta"].get("price_sheet") or "não encontrada"),
        ("SHA-256 da referência", result["meta"]["sha256"]), ("Gerado em", result["created_at"]),
        ("Validação", "Erro médio absoluto de pontos internos retirados, dividido pela média observada; não é probabilidade de acerto."),
        ("Critério de revisão", "Erro de validação acima de 20%: revisar. Critério do programa; não é norma SINAPI."),
        ("Dados", "Cada coeficiente mantém a unidade do componente por unidade do serviço. Os valores não são somados entre unidades diferentes.")]
    if result.get("calculation_type") == "rule_of_three":
        metadata = [(key, "Não se aplica: uma referência não permite validar tendência." if key == "Validação"
                     else "Proporcionalidade assumida; revisão técnica obrigatória." if key == "Critério de revisão" else value)
                    for key, value in metadata]
        metadata += [("Fator dimensional", result["factor"]), ("Cálculo do fator", result["factor_formula"]),
                     ("Aplicação", "Somente nos componentes marcados; consulte as colunas de fator por recurso.")]
    elif result.get("calculation_type") == "manual_trend":
        metadata += [("Entrada manual", "Dimensões e comparabilidade informadas pelo usuário. Coeficientes podem ter sido digitados ou editados."),
                     ("Referência no alvo", result.get("target_reference") or "Alvo não coincide com ponto informado")]
    metadata += [("Observação", w) for w in result["warnings"]]
    for row in metadata:
        summary.append(row)
    sheet = book.create_sheet("Coeficientes")
    headers = ["Tipo", "Código do insumo na composição de origem", "Descrição na composição de origem",
                  "Código do insumo na nova composição", "Descrição do insumo na nova composição",
                  "Unidade componente", "Coef. origem", "Coef. alvo", "Variação %", "Situação",
                  "Validação NMAE %", "Pontos de validação", "Verificação de insumo SINAPI"]
    proportional = result.get("calculation_type") == "rule_of_three"
    if proportional:
        headers += ["Aplicar fator?", "Fator aplicado", "Cálculo por componente"]
    sheet.append(headers)
    for row in result["rows"]:
        values = [row["kind"], row["source_resource"], row["description"], row["target_resource"], row["target_description"], row["unit"], row["source_coefficient"], row["coefficient"], row["change_pct"], row["status"], row["validation"]["nmae_pct"], row["validation"]["n"], row["supply_status"]]
        if proportional:
            values += ["Sim" if row["proportional_applied"] else "Manter", row["applied_factor"], row["calculation"]]
        sheet.append(values)
    refs = book.create_sheet("Composições de referência")
    manual = result.get("calculation_type") == "manual_trend"
    refs.append(["Código/nome" if manual else "Código", "Grandeza informada" if manual else "Geometria interpretada", "Dimensões informadas", "Índice da curva", "Descrição da referência" if manual else "Descrição SINAPI original", "Linha na aba Analítico", "Compatibilidade declarada" if manual else "Conjunto compatível"])
    for ref in result["references"]:
        refs.append([ref["code"], ref["dimension_text"], ref.get("reported_dimensions"), ref["x"], ref["description"], ref["row"], "Sim" if ref["used"] else "Não"])
    points = book.create_sheet("Pontos e variações")
    points.append(["Recurso de origem", "Descrição", "Referência informada" if manual else "Composição SINAPI", "Grandeza" if manual else "Geometria", "Índice da curva", "Código recurso", "Coeficiente informado" if manual else "Coeficiente observado", "Linha Analítico", "Variação desde ponto anterior %", "Inclinação do intervalo"])
    for row in result["rows"]:
        previous = None
        for point in row["points"]:
            y = point["coefficient"]
            valid = previous and previous["coefficient"] is not None and y is not None
            change = (y / previous["coefficient"] - 1) * 100 if valid and previous["coefficient"] else None
            slope = (y - previous["coefficient"]) / (point["x"] - previous["x"]) if valid and point["x"] != previous["x"] else None
            points.append([row["source_resource"], row["description"], point["composition"], point["dimension_text"], point["x"], point["resource"], y, point["row"], change, slope])
            previous = point
    supplies = book.create_sheet("Insumos SINAPI PE-SP")
    supplies.append(["Componente da CPU", "Verificação", "Código do insumo", "Descrição do insumo", "Unidade", "Preço PE", "Preço SP", "Origem do preço", "Aba", "Linha"])
    for row in result["rows"]:
        if row["supply_matches"]:
            for item in row["supply_matches"]:
                supplies.append([row["source_resource"], row["supply_status"], item["code"], item["description"], item["unit"], item.get("PE"), item.get("SP"), item.get("origin"), item.get("sheet"), item.get("row")])
        elif row["supply_checked"]:
            supplies.append([row["source_resource"], row["supply_status"], "", "", "", None, None, "", result["meta"].get("price_sheet"), ""])
    for ws in book:
        ws.freeze_panes = "A2"
        if ws != summary:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="173D37")
            cell.font = Font(color="FFFFFF", bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    # A descrição do arquivo é texto, nunca uma fórmula executável.
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = "0.00000000"
        for col in range(1, ws.max_column + 1):
            maximum = max((len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 150) + 1)), default=15)
            ws.column_dimensions[get_column_letter(col)].width = min(75, max(15, maximum + 2))
    path = Path(path)
    temporary = path.with_name(path.stem + ".tmp.xlsx")
    try:
        book.save(temporary)
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()
    return path
