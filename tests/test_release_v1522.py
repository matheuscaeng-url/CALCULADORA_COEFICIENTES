"""Cobertura real de eletrocalhas, altura editável e fluxos da v1.5.22."""
import copy
import inspect
from itertools import combinations
import math
from pathlib import Path
import queue
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP))

from coeficientes_app import CalculatorApp
from coeficientes_core import (
    AnalysisError, SCOPE_REVISION, aligned_resources, build_families, calculate,
    cable_tray_geometry, cable_tray_signature, connection_geometry, export_report,
    geometry_index, import_sinapi, is_automatic_scope, is_cable_tray,
    is_pipe_tee, is_tube_or_fitting, load_database, save_database,
)
from tendencia_manual_core import calculate_manual_trend, series_from_sinapi


class CableTrayReleaseTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.db = load_database(APP / 'data/sinapi_base.json.gz')
        cls.full = load_database(APP / 'data/sinapi_completa.json.gz')
        cls.families = build_families(cls.db)
        cls.electric = [f for f in cls.families if f['discipline'] == 'Elétrica']

    def family(self, code):
        return next(f for f in self.electric if any(r['code'] == code for r in f['refs']))

    def test_all_actual_compositions_have_exactly_one_electric_family(self):
        expected = {str(code) for code in range(97236, 97326)} | {'96562', '96563'}
        refs = [r['code'] for f in self.electric for r in f['refs']]
        self.assertEqual(set(refs), expected)
        self.assertEqual(len(refs), len(expected))
        self.assertEqual(sorted(len(f['refs']) for f in self.electric), [2, 12, 13, 13, 13, 13, 13, 13])
        self.assertTrue(all(f['has_trend'] and not f['tee'] for f in self.electric))
        self.assertEqual(len(self.db['compositions']), 2300)
        self.assertEqual(len(self.full['compositions']), 10547)
        self.assertEqual(len([f for f in self.families if not f.get('cable_tray')]), 609)

    def test_cables_and_pipe_tees_do_not_cross_scope(self):
        for code in ['101884', '101886', '101888']:
            desc = self.full['compositions'][code]['description']
            self.assertFalse(is_cable_tray(desc))
            self.assertFalse(is_automatic_scope(desc))
            self.assertNotIn(code, self.db['compositions'])
        desc = self.db['compositions']['97313']['description']
        self.assertTrue(is_cable_tray(desc))
        self.assertFalse(is_tube_or_fitting(desc))
        self.assertFalse(is_pipe_tee(desc))
        self.assertIsNone(connection_geometry(desc))
        self.assertTrue(is_pipe_tee('TÊ DE REDUÇÃO, PVC, SOLDÁVEL, DN 40 X 32 MM'))

    def test_width_equal_to_height_still_has_two_independent_axes(self):
        for code in ['97236', '97249', '97274', '97275', '97276', '97313']:
            family = self.family(code)
            self.assertEqual(family['dimension_count'], 2)
            self.assertEqual(family['axis_labels'], ['Largura', 'Altura'])
            self.assertEqual(family['height_axis'], 1)
            self.assertEqual(next(r['dims'] for r in family['refs'] if r['code'] == code), [50, 50])
        geometry = cable_tray_geometry(self.db['supplies']['41549']['description'])
        self.assertEqual(geometry['dims'], [50, 50])
        self.assertEqual(geometry['height_mm'], 50)

    def test_reduction_has_two_widths_and_height_in_both_descriptions(self):
        for code, supply, widths in [('97252', '43966', [100, 75]), ('97273', '44048', [700, 600])]:
            for desc in [self.db['compositions'][code]['description'], self.db['supplies'][supply]['description']]:
                geometry = cable_tray_geometry(desc)
                self.assertEqual(geometry['dims'], widths + [50])
                self.assertEqual(geometry['height_mm'], 50)
                self.assertEqual(geometry['kind'], 'tray_reduction')
        family = self.family('97252')
        self.assertEqual(family['dimension_count'], 3)
        self.assertEqual(family['axis_labels'], ['Largura maior', 'Largura menor', 'Altura'])
        self.assertAlmostEqual(family['refs'][0]['x'], math.sqrt((75**2 + 50**2 + 50**2) / 3))

    def test_different_heights_join_but_technical_specs_stay_separate(self):
        source = copy.deepcopy(self.db['compositions']['97236'])
        data = {'compositions': {'A': source}}
        for code, before, after in [('B', 'ALTURA 50MM', 'ALTURA 100MM'),
                                    ('C', 'AÇO GALVANIZADO', 'AÇO INOXIDÁVEL')]:
            item = copy.deepcopy(source)
            item['code'] = code
            item['description'] = item['description'].replace(before, after)
            data['compositions'][code] = item
        families = build_families(data)
        self.assertEqual(len(families), 2)
        shared = next(f for f in families if len(f['refs']) == 2)
        self.assertEqual({r['code'] for r in shared['refs']}, {'A', 'B'})
        low = self.db['supplies']['41549']['description']
        high = self.db['supplies']['44047']['description']
        self.assertEqual(cable_tray_signature(low, align_gauge=True), cable_tray_signature(high, align_gauge=True))
        self.assertNotEqual(cable_tray_signature(low), cable_tray_signature(high))

    def test_all_references_align_including_gauges_and_auxiliary_supports(self):
        for family in self.electric:
            maps = aligned_resources(self.db, family['refs'], family)
            self.assertTrue(all(set(m) == set(maps[0]) for m in maps), family['title'])
        straight = self.family('97236')
        self.assertEqual(len(aligned_resources(self.db, straight['refs'], straight)[0]), 5)

    def test_all_92_official_targets_return_exact_original_components(self):
        for family in self.electric:
            for ref in family['refs']:
                result = calculate(self.db, family, family['refs'][0]['code'], ref['dims'])
                self.assertEqual(result['target_code'], ref['code'])
                expected = {c['code']: (c['coefficient'], c['description'])
                            for c in self.db['compositions'][ref['code']]['components']}
                found = {r['target_resource']: (r['coefficient'], r['target_description']) for r in result['rows']}
                self.assertEqual(found, expected, ref['code'])

    def test_all_eight_families_predict_using_every_reference(self):
        for family in self.electric:
            target = ([600, 65] if family['geometry_kind'] == 'tray_support'
                      else [175, 140, 50] if family['geometry_kind'] == 'tray_reduction' else [175, 50])
            result = calculate(self.db, family, family['refs'][0]['code'], target)
            self.assertIsNone(result['target_code'])
            self.assertTrue(all(math.isfinite(r['coefficient']) and r['coefficient'] >= 0 for r in result['rows']))
            self.assertTrue(all(len(r['points']) == len(family['refs']) for r in result['rows']))
            self.assertNotIn('diâmetro equivalente', result['method'])

    def test_reduction_both_widths_affect_prediction(self):
        family = self.family('97252')
        a = calculate(self.db, family, '97252', [175, 140, 50])
        b = calculate(self.db, family, '97252', [175, 145, 50])
        labor = lambda result: next(r['coefficient'] for r in result['rows'] if r['source_resource'] == '88264')
        self.assertNotEqual(labor(a), labor(b))
        self.assertIn('índice RMS', a['justification'])
        with self.assertRaisesRegex(AnalysisError, 'largura maior'):
            calculate(self.db, family, '97252', [100, 150, 50])

    def test_extrapolation_still_requires_explicit_opt_in(self):
        family = self.family('97236')
        with self.assertRaisesRegex(AnalysisError, 'extrapolação'):
            calculate(self.db, family, '97236', [900, 50])
        result = calculate(self.db, family, '97236', [900, 50], True)
        self.assertIn('extrapolação', result['mode'])

    def test_support_dimensions_are_coupled_and_unit_conversion_is_correct(self):
        family = self.family('96562')
        self.assertEqual(family['axis_units'], ['MM', 'CM'])
        self.assertEqual(family['profile_offset_mm'], 50)
        result = calculate(self.db, family, '96562', [600, 65])
        x1, x2, target = [geometry_index(d, ['MM', 'CM']) for d in ([400, 45], [800, 85], [600, 65])]
        expected = 0.5 + (0.9444 - 0.5) * (target - x1) / (x2 - x1)
        profile = next(r for r in result['rows'] if r['source_resource'] == '39029')
        self.assertAlmostEqual(profile['coefficient'], expected)
        self.assertEqual(profile['supply_matches'][0]['code'], '39029')
        with self.assertRaisesRegex(AnalysisError, 'mais 50 mm'):
            calculate(self.db, family, '96562', [600, 45])

    def test_estimated_material_does_not_inherit_gauge_or_invent_aux_support(self):
        result = calculate(self.db, self.family('97236'), '97236', [175, 50])
        material = next(r for r in result['rows'] if r['source_resource'] == '41549')
        self.assertIn('175 MM', material['target_description'])
        self.assertIn('50 MM', material['target_description'])
        self.assertNotIn('#22', material['target_description'])
        self.assertIn('ESPESSURA A DEFINIR', material['target_description'])
        support = next(r for r in result['rows'] if r['source_resource'] == '96562')
        self.assertIn('LARGURA 400 MM', support['target_description'])
        self.assertNotIn('LARGURA 175', support['target_description'])
        self.assertIn('CONFERIR O SUPORTE', support['target_description'])
        self.assertIn('DEFINIR', support['target_resource'])

    def test_all_electrical_materials_and_unchanged_hardware_have_prices_from_isd(self):
        codes = {c['code'] for f in self.electric for ref in f['refs']
                 for c in self.db['compositions'][ref['code']]['components'] if c['kind'] == 'INSUMO'}
        self.assertEqual(len(codes), 99)
        self.assertTrue(codes <= set(self.db['supplies']))
        self.assertEqual(len(self.db['supplies']), 1366)
        result = calculate(self.db, self.family('97249'), '97249', [175, 50])
        for code in ['45693', '44027', '43965', '4330']:
            row = next(r for r in result['rows'] if r['source_resource'] == code)
            self.assertEqual(row['supply_matches'], [self.db['supplies'][code]])
            self.assertEqual(row['target_resource'], code)

    def test_reimport_preserves_source_compositions_and_extended_scope(self):
        imported = import_sinapi(APP / 'REFERENCIA_ORIGINAL/SINAPI_Referencia_2026_08.xlsx')
        self.assertEqual(imported['compositions'], self.db['compositions'])
        self.assertEqual(imported['supplies'], self.db['supplies'])
        self.assertEqual(imported['meta']['scope_revision'], SCOPE_REVISION)

    def test_electric_result_exports_widths_and_height(self):
        from openpyxl import load_workbook
        result = calculate(self.db, self.family('97252'), '97252', [175, 140, 50])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'reducao.xlsx'
            export_report(result, path)
            book = load_workbook(path, data_only=True)
            summary = dict(book['Leia primeiro'].iter_rows(values_only=True))
            self.assertEqual(summary['Dimensão alvo'], '175 mm × 140 mm × 50 mm')
            self.assertIn('índice RMS', summary['Método'])
            self.assertTrue(any('Alturas observadas' in str(row) and '50 mm' in str(row)
                                for row in book['Leia primeiro'].values))
            book.close()

    def test_ui_uses_width_and_profile_labels(self):
        for code in ['97313', '97252', '96562']:
            family = self.family(code)
            app = SimpleNamespace(invalidate=Mock(), family_title=Mock(), source_box=Mock(), status=Mock(),
                                  target_labels=[Mock() for _ in range(3)], target_entries=[Mock() for _ in range(3)],
                                  targets=[Mock() for _ in range(3)])
            CalculatorApp.activate_family(app, family)
            for i, expected in enumerate(family['axis_labels']):
                self.assertIn(expected, app.target_labels[i].configure.call_args.kwargs['text'])
                self.assertNotIn('Boca', app.target_labels[i].configure.call_args.kwargs['text'])

    def test_height_100_requires_extrapolation_even_inside_rms_range_for_every_fitting(self):
        for family in self.electric:
            if family['geometry_kind'] == 'tray_support':
                continue
            target = [175, 140, 100] if family['geometry_kind'] == 'tray_reduction' else [175, 100]
            code = family['refs'][0]['code']
            with self.subTest(family=family['title']):
                with self.assertRaisesRegex(AnalysisError, 'altura.*extrapolação'):
                    calculate(self.db, family, code, target)
                result = calculate(self.db, family, code, target, True)
                self.assertIsNone(result['target_code'])
                self.assertIn('Extrapolação dimensional via PCHIP', result['method'])
                self.assertTrue(result['range'][0] < result['target'] < result['range'][1])
                self.assertTrue(all(row['coefficient'] is not None and row['coefficient'] >= 0
                                    for row in result['rows']))
                self.assertIn('altura sem referência observada', result['justification'])
                self.assertTrue(any('ALTURA SEM REFERÊNCIA OBSERVADA' in warning and '100 mm' in warning
                                    for warning in result['warnings']))

    def test_changing_height_changes_labor_and_does_not_claim_existing_width_code(self):
        family = self.family('97236')
        official = calculate(self.db, family, '97236', [100, 50])
        adapted = calculate(self.db, family, '97236', [100, 100], True)
        self.assertEqual(official['target_code'], '97238')
        self.assertIsNone(adapted['target_code'])
        labor = lambda result: next(row['coefficient'] for row in result['rows']
                                    if row['source_resource'] == '88264')
        self.assertNotEqual(labor(official), labor(adapted))
        material = next(row for row in adapted['rows'] if row['source_resource'] == '41549')
        self.assertEqual(cable_tray_geometry(material['target_description'])['dims'], [100, 100])
        self.assertEqual(material['supply_matches'], [])
        self.assertIn('DEFINIR', material['target_resource'])

    def test_reduction_updates_all_three_material_dimensions(self):
        family = self.family('97252')
        result = calculate(self.db, family, '97252', [175, 140, 100], True)
        material = next(row for row in result['rows'] if row['source_resource'] == '43966')
        self.assertEqual(cable_tray_geometry(material['target_description'])['dims'], [175, 140, 100])
        self.assertEqual(result['target_dimension_text'], '175 mm × 140 mm × 100 mm')
        self.assertEqual(material['supply_matches'], [])
        self.assertIn('97252', result['justification'])
        self.assertIn('100 mm', result['justification'])

    def test_height_dependent_tala_changes_description_code_and_price_status(self):
        family = self.family('97249')
        original = calculate(self.db, family, '97249', [175, 50])
        adapted = calculate(self.db, family, '97249', [175, 100], True)
        old = next(row for row in original['rows'] if row['source_resource'] == '44027')
        new = next(row for row in adapted['rows'] if row['source_resource'] == '44027')
        self.assertEqual(old['target_resource'], '44027')
        self.assertEqual(old['supply_matches'], [self.db['supplies']['44027']])
        self.assertIn('DEFINIR', new['target_resource'])
        self.assertIn('ALTURA 100 MM', new['target_description'])
        self.assertNotIn('ALTURA 50 MM', new['target_description'])
        self.assertEqual(new['supply_matches'], [])
        for code in ['45693', '43965', '4330']:
            row = next(row for row in adapted['rows'] if row['source_resource'] == code)
            self.assertEqual(row['target_resource'], code)
            self.assertEqual(row['supply_matches'], [self.db['supplies'][code]])

    def test_target_height_supply_is_found_only_when_it_really_exists(self):
        data = copy.deepcopy(self.db)
        new_supply = copy.deepcopy(data['supplies']['44027'])
        new_supply.update(code='TEST100', PE=123.45, SP=234.56)
        new_supply['description'] = new_supply['description'].replace('ALTURA 50 MM', 'ALTURA 100 MM')
        data['supplies']['TEST100'] = new_supply
        result = calculate(data, self.family('97249'), '97249', [175, 100], True)
        row = next(row for row in result['rows'] if row['source_resource'] == '44027')
        self.assertEqual(row['target_resource'], 'TEST100')
        self.assertEqual(row['target_description'], new_supply['description'])
        self.assertEqual(row['supply_matches'], [new_supply])
        self.assertIn('PE e SP', row['supply_status'])

    def test_future_reference_with_another_height_is_grouped_and_consulted_directly(self):
        data = copy.deepcopy(self.db)
        item = copy.deepcopy(data['compositions']['97238'])
        item['code'] = 'TEST100'
        item['description'] = item['description'].replace('ALTURA 50MM', 'ALTURA 100MM')
        material = next(c for c in item['components'] if is_cable_tray(c['description'])
                        and cable_tray_geometry(c['description'])['kind'] == 'tray')
        material['code'] = 'TEST_MATERIAL100'
        material['description'] = material['description'].replace('X 50 MM', 'X 100 MM')
        data['compositions']['TEST100'] = item
        family = next(f for f in build_families(data) if any(r['code'] == 'TEST100' for r in f['refs']))
        self.assertEqual(len(family['refs']), 14)
        result = calculate(data, family, '97236', [100, 100])
        self.assertEqual(result['target_code'], 'TEST100')
        self.assertEqual(result['method'], 'Consulta direta; sem estimativa')
        expected = {c['code']: (c['coefficient'], c['description']) for c in item['components']}
        found = {r['target_resource']: (r['coefficient'], r['target_description']) for r in result['rows']}
        self.assertEqual(found, expected)

    def test_height_adaptation_export_contains_target_and_limitation(self):
        from openpyxl import load_workbook
        result = calculate(self.db, self.family('97236'), '97236', [175, 100], True)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'eletrocalha_altura100.xlsx'
            export_report(result, path)
            book = load_workbook(path, data_only=True)
            try:
                summary = dict(book['Leia primeiro'].iter_rows(values_only=True))
                self.assertEqual(summary['Dimensão alvo'], '175 mm × 100 mm')
                self.assertIn('altura sem referência observada', str(summary))
                self.assertTrue(any('ALTURA SEM REFERÊNCIA OBSERVADA' in str(row)
                                    for row in book['Leia primeiro'].values))
            finally:
                book.close()

    def test_height_is_required_and_positive(self):
        family = self.family('97236')
        for target in ([175], [175, 0], [175, -100]):
            with self.subTest(target=target), self.assertRaises(AnalysisError):
                calculate(self.db, family, '97236', target, True)


class ManualAlignmentReleaseTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.full = load_database(APP / 'data/sinapi_completa.json.gz')
        cls.electric = [f for f in build_families(load_database(APP / 'data/sinapi_base.json.gz'))
                        if f.get('cable_tray')]

    def load_pair(self, data, codes=('A', 'B')):
        refs = [{'code': code, 'dims': [10 + i * 10]} for i, code in enumerate(codes)]
        return refs, series_from_sinapi(data, refs)[0]

    def synthetic(self, descriptions, *, units=('UN', 'UN'), kinds=('INSUMO', 'INSUMO')):
        return {'compositions': {
            code: {'code': code, 'description': 'Referência ' + code, 'unit': 'UN',
                   'components': [{'code': 'ITEM_' + code, 'description': description,
                                   'unit': units[i], 'kind': kinds[i], 'coefficient': i + 1, 'row': i + 1}]}
            for i, (code, description) in enumerate(zip(('A', 'B'), descriptions))}}

    def test_reported_pairs_have_five_five_and_seven_real_rows(self):
        for a, b, expected_count in [('97236', '97238', 5), ('97238', '97248', 5),
                                     ('97249', '97251', 7)]:
            with self.subTest(codes=(a, b)):
                refs, series = self.load_pair(self.full, (a, b))
                self.assertEqual(len(series), expected_count)
                self.assertTrue(all(r['alignment_complete'] for r in series))
                for column, code in enumerate((a, b)):
                    components = self.full['compositions'][code]['components']
                    recovered = sorted(series, key=lambda r: r['orders'][column])
                    self.assertEqual([(r['codes'][column], r['coefficients'][column], r['rows'][column])
                                      for r in recovered],
                                     [(c['code'], c['coefficient'], c['row']) for c in components])
                result = calculate_manual_trend(refs, series, 15)
                self.assertEqual(len(result['rows']), expected_count)
                self.assertTrue(all(r['coefficient'] is not None for r in result['rows']))

    def test_every_pair_in_the_eight_tray_families_aligns_without_missing_cells(self):
        for family in self.electric:
            for a, b in combinations(family['refs'], 2):
                with self.subTest(codes=(a['code'], b['code'])):
                    _, series = self.load_pair(self.full, (a['code'], b['code']))
                    self.assertEqual(len(series), len(self.full['compositions'][a['code']]['components']))
                    self.assertTrue(all(row['alignment_complete'] for row in series))

    def test_generic_sequences_typography_and_fractions_align(self):
        examples = [
            ('CAIXA PVC, 20 X30MM, CLASSE A, NBR 12345',
             'CAIXA PVC - 40X60MM, CLASSE A, NBR 12345'),
            ('CAIXA PVC 20 MM X 30 MM X 40 MM', 'CAIXA PVC 40 × 50 × 60MM'),
            ('PEÇA, COMPRIMENTO DE 1½″', 'PEÇA - COMPRIMENTO 2.1/2"'),
            ('PEÇA 1/2" X 3/4"', 'PEÇA 1” × 1 1/2“'),
            ('PERFIL, LARGURA DE 20 MM, ALTURA = 30 MM',
             'PERFIL - LARGURA 40MM - ALTURA 60MM'),
        ]
        for descriptions in examples:
            with self.subTest(descriptions=descriptions):
                _, series = self.load_pair(self.synthetic(descriptions))
                self.assertEqual(len(series), 1)
                self.assertEqual(series[0]['codes'], ['ITEM_A', 'ITEM_B'])
                self.assertEqual(series[0]['coefficients'], [1, 2])

    def test_material_joint_angle_norm_class_kind_unit_stay_distinct_but_height_aligns(self):
        base = 'JOELHO PVC SOLDAVEL 45 GRAUS, DN 20 MM, CLASSE 15, NBR 5648'
        for before, after in [('PVC', 'CPVC'), ('SOLDAVEL', 'ROSCAVEL'), ('45', '90'),
                              ('NBR 5648', 'NBR 15884'), ('CLASSE 15', 'CLASSE 20')]:
            other = base.replace('20 MM', '40 MM').replace(before, after)
            with self.subTest(change=after):
                _, series = self.load_pair(self.synthetic((base, other)))
                self.assertEqual(len(series), 2)
                self.assertEqual([r['coefficients'] for r in series], [[1, ''], ['', 2]])
        for options in ({'units': ('UN', 'M')}, {'kinds': ('INSUMO', 'COMPOSICAO')}):
            _, series = self.load_pair(self.synthetic((base, base), **options))
            self.assertEqual(len(series), 2)
        tray = self.full['compositions']['97236']['description']
        _, series = self.load_pair(self.synthetic((tray, tray.replace('ALTURA 50MM', 'ALTURA 100MM'))))
        self.assertEqual(len(series), 1)
        self.assertEqual(series[0]['coefficients'], [1, 2])

    def test_repeated_components_extra_resource_and_zero_are_preserved_one_to_one(self):
        data = self.synthetic(('PEÇA PVC 20 X 30 MM', 'PEÇA PVC 40 X 60 MM'))
        for column, code in enumerate(('A', 'B')):
            components = data['compositions'][code]['components']
            components[0]['coefficient'] = 0
            repeated = copy.deepcopy(components[0])
            repeated['coefficient'] = column + 5
            repeated['row'] += 10
            components.append(repeated)
        data['compositions']['B']['components'].append({
            'code': 'EXTRA', 'description': 'FITA ISOLANTE', 'unit': 'UN', 'kind': 'INSUMO',
            'coefficient': 3, 'row': 99})
        _, series = self.load_pair(data)
        self.assertEqual(len(series), 3)
        self.assertEqual([r['coefficients'] for r in series], [[0, 0], [5, 6], ['', 3]])
        self.assertEqual([r['orders'] for r in series], [[0, 0], [1, 1], [None, 2]])
        self.assertEqual([r['rows'] for r in series], [[1, 2], [11, 12], ['', 99]])
        self.assertFalse(series[-1]['alignment_complete'])


class MigrationAndUiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base = load_database(APP / 'data/sinapi_base.json.gz')
        cls.full = load_database(APP / 'data/sinapi_completa.json.gz')

    def legacy(self, full=False):
        data = copy.deepcopy(self.full if full else self.base)
        data['meta'].pop('scope_revision', None)
        data['supplies'] = {code: v for code, v in data['supplies'].items() if is_tube_or_fitting(v['description'])}
        if not full:
            data['compositions'] = {code: v for code, v in data['compositions'].items() if is_tube_or_fitting(v['description'])}
        return data

    def test_old_local_full_base_gets_matching_supplies_and_preserves_custom_values(self):
        active, full = self.legacy(), self.legacy(full=True)
        active['compositions']['94464']['components'][0]['coefficient'] = 9.8765
        full['supplies']['7694']['PE'] = 123.45
        with tempfile.TemporaryDirectory() as directory, patch('coeficientes_app.USER_DIR', Path(directory)):
            save_database(full, Path(directory) / 'base_completa.json.gz')
            matched = CalculatorApp.matching_full_database(active)
            upgraded = CalculatorApp.upgrade_automatic_database(active, matched)
            self.assertEqual(upgraded['meta']['period'], active['meta']['period'])
            self.assertEqual(upgraded['meta']['sha256'], active['meta']['sha256'])
            self.assertEqual(upgraded['compositions']['94464']['components'][0]['coefficient'], 9.8765)
            self.assertEqual(matched['supplies']['7694']['PE'], 123.45)
            self.assertEqual(len(upgraded['compositions']), 2300)
            self.assertEqual(len(upgraded['supplies']), 1366)
            self.assertEqual(CalculatorApp.upgrade_automatic_database(upgraded, matched), upgraded)

    def test_custom_period_is_extended_only_from_its_own_full_base(self):
        active, full = self.legacy(), self.legacy(full=True)
        for data in (active, full):
            data['meta'].update(period='01/2025', sha256='custom-reference')
        with tempfile.TemporaryDirectory() as directory, patch('coeficientes_app.USER_DIR', Path(directory)):
            self.assertIsNone(CalculatorApp.matching_full_database(active))
            self.assertIs(CalculatorApp.upgrade_automatic_database(active, self.full), active)
            save_database(full, Path(directory) / 'base_completa.json.gz')
            matched = CalculatorApp.matching_full_database(active)
            upgraded = CalculatorApp.upgrade_automatic_database(active, matched)
            self.assertEqual(upgraded['meta']['period'], '01/2025')
            self.assertEqual(len(upgraded['compositions']), 2300)
            self.assertEqual(len(upgraded['supplies']), 1267)

    def test_initial_loading_migrates_and_persists_old_active_database(self):
        with tempfile.TemporaryDirectory() as directory, patch('coeficientes_app.USER_DIR', Path(directory)):
            active = Path(directory) / 'base_ativa.json.gz'
            save_database(self.legacy(), active)
            app = SimpleNamespace(matching_full_database=CalculatorApp.matching_full_database,
                                  upgrade_automatic_database=CalculatorApp.upgrade_automatic_database,
                                  set_database=Mock())
            app.start_job = lambda work, callback, message: callback(work())
            CalculatorApp.load_initial(app)
            payload = app.set_database.call_args.args[0]
            self.assertEqual(len(payload[0]['compositions']), 2300)
            self.assertEqual(len(load_database(active)['supplies']), 1366)

    def test_removed_control_has_no_dialog_or_busy_state_dependency(self):
        self.assertFalse(hasattr(CalculatorApp, 'manual_dialog'))
        self.assertNotIn('Selecionar referências por código', inspect.getsource(CalculatorApp))
        app = SimpleNamespace(working=False, status=Mock(), progress=Mock(), events=queue.Queue(),
                              rule_tab=Mock(), manual_tab=Mock(), result=None, after=Mock(), poll=Mock(),
                              **{name: Mock() for name in ('import_button', 'restore_button', 'calc_button',
                                                         'export_button', 'copy_button')})
        with patch('coeficientes_app.threading.Thread'):
            CalculatorApp.start_job(app, lambda: None, Mock(), 'Teste')
        self.assertTrue(app.working)
        callback = Mock()
        app.events.put(('done', callback, 'resultado'))
        CalculatorApp.poll(app)
        self.assertFalse(app.working)
        callback.assert_called_once_with('resultado')


if __name__ == '__main__':
    unittest.main(verbosity=2)
