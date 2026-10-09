import random
import sys
import tempfile
import unittest
from pathlib import Path

BASE = Path(__file__).resolve().parents[1]
APP = BASE if (BASE / "coeficientes_core.py").exists() else BASE / "PROCESSO_3_COEFICIENTES"
sys.path.insert(0, str(APP))

from coeficientes_core import (AnalysisError, build_families, calculate, export_report,
    connection_geometry, dimension_number, dimension_range, format_geometry,
    format_reference_dimensions, is_automatic_scope, is_tube_or_fitting, load_database, manual_family,
    measures, nominal_dimensions_equal, norm, number, pchip, tee_geometry,
    tee_signature, validation)


def synthetic(xs=(20, 25, 40, 60), ys=(.1, .15, .28, .55)):
    comps = {}
    for i, (x, y) in enumerate(zip(xs, ys)):
        comps[str(i)] = {"code": str(i), "description": f"TUBO PVC DN {x} MM AF_01/2020",
            "group": "Tubos", "unit": "M", "status": "COM CUSTO", "row": 10 + i * 3,
            "components": [{"code": "88267", "description": "ENCANADOR COM ENCARGOS COMPLEMENTARES",
                "kind": "COMPOSICAO", "unit": "H", "coefficient": y, "row": 11 + i * 3}]}
    return {"schema": 2, "meta": {"period": "08/2026", "filename": "fixture", "sheet": "Analítico",
            "price_sheet": "ISD", "sha256": "test"}, "compositions": comps, "supplies": {}}


class NumericalTests(unittest.TestCase):
    def test_reference_points_and_bounds(self):
        xs, ys = [20, 25, 40, 60], [.1, .15, .28, .55]
        for x, y in zip(xs, ys):
            self.assertAlmostEqual(pchip(xs, ys, x), y)
        for a, b, ya, yb in zip(xs, xs[1:], ys, ys[1:]):
            for i in range(101):
                self.assertTrue(ya - 1e-12 <= pchip(xs, ys, a + (b-a)*i/100) <= yb + 1e-12)

    def test_matches_independent_scipy_implementation(self):
        try:
            from scipy.interpolate import PchipInterpolator
        except ImportError:
            self.skipTest("SciPy é usada somente como verificação independente")
        generator = random.Random(8421)
        for _ in range(30):
            xs = sorted(generator.sample(range(1, 400), 7))
            ys = [generator.random() * 3 for _ in xs]
            reference = PchipInterpolator(xs, ys)
            for i in range(21):
                x = xs[0] + (xs[-1]-xs[0])*i/20
                self.assertAlmostEqual(pchip(xs, ys, x), float(reference(x)), places=11)

    def test_two_points_are_allowed_with_warning(self):
        data = synthetic((20, 40), (.1, .3))
        family = build_families(data)[0]
        result = calculate(data, family, "0", 30)
        self.assertAlmostEqual(result["rows"][0]["coefficient"], .2)
        self.assertIn("2 pontos", " ".join(result["warnings"]))
        self.assertIn("maior possibilidade de erro", result["rows"][0]["validation"]["label"])

    def test_extrapolation_has_no_twenty_percent_cap(self):
        data = synthetic()
        family = build_families(data)[0]
        with self.assertRaises(AnalysisError):
            calculate(data, family, "0", 200)
        result = calculate(data, family, "0", 200, True)
        self.assertIn("extrapolação", result["mode"])
        self.assertIn("SEM LIMITE", " ".join(result["warnings"]))

    def test_scope_only_tubes_and_pipe_fittings(self):
        self.assertTrue(is_tube_or_fitting("TUBO, PVC, SOLDÁVEL, DN 25 MM, INSTALADO EM PRUMADA DE ÁGUA"))
        self.assertTrue(is_tube_or_fitting("TÊ DE REDUÇÃO, PVC, SOLDÁVEL, DN 40 X 32 MM"))
        self.assertFalse(is_tube_or_fitting("CABO DE COBRE FLEXÍVEL 2,5 MM2"))
        self.assertFalse(is_tube_or_fitting("FABRICAÇÃO DE TÊ PARA DUTO DE AR EM CHAPA GALVANIZADA"))
        self.assertFalse(is_tube_or_fitting("ELETRODUTO RÍGIDO ROSCÁVEL PVC DN 25 MM"))

    def test_tee_one_two_and_three_dimensions(self):
        one = tee_geometry("TÊ, PVC, SOLDÁVEL, DN 25 MM, INSTALADO EM PRUMADA DE ÁGUA")
        two = tee_geometry("TÊ DE REDUÇÃO, PVC, SOLDÁVEL, DN 40 X 32 MM, INSTALADO EM PRUMADA DE ÁGUA")
        three = tee_geometry("TÊ DE REDUÇÃO PARA PEX, DN 25 X 20 X 20 MM")
        self.assertEqual(one["dims"], [25, 25, 25])
        self.assertEqual(two["dims"], [40, 40, 32])
        self.assertEqual(three["dims"], [25, 20, 20])
        self.assertNotEqual(tee_signature("TÊ, PVC, SOLDÁVEL, DN 40 MM, INSTALADO EM PRUMADA DE ÁGUA"),
                            tee_signature("TÊ DE REDUÇÃO, PVC, SOLDÁVEL, DN 40 X 32 MM, INSTALADO EM PRUMADA DE ÁGUA"))

    def test_manual_family_accepts_two_and_three_axes(self):
        data = synthetic((20, 40), (.1, .3))
        family = manual_family(data, [("0", "20", "20", "20"), ("1", "40", "40", "32")], "MM")
        self.assertEqual(family["dimension_count"], 3)
        self.assertEqual(len(family["refs"]), 2)

    def test_numbers_fractional_inches_and_validation(self):
        self.assertEqual(number("1.234,56"), 1234.56)
        self.assertEqual(measures('LUVA 1 1/4"')[0]["value"], 1.25)
        self.assertEqual(dimension_number('1 1/8"'), 1.125)
        self.assertEqual(format_geometry([1.125, 1.625], ["POL", "POL"]), '1 1/8" × 1 5/8"')
        self.assertEqual(validation([1, 2], [1, 2])["n"], 0)
        with self.assertRaises(ValueError):
            number("nan")

    def test_bare_dn_and_nominal_pipe_equivalence(self):
        axis = measures("ASSENTAMENTO DE CONEXÃO DE FERRO FUNDIDO, DN 150, JUNTA ELÁSTICA")[0]
        self.assertEqual((axis["value"], axis["unit"]), (150, "MM"))
        self.assertTrue(nominal_dimensions_equal(80, "MM", 3, "POL"))
        self.assertFalse(nominal_dimensions_equal(75, "MM", 3, "POL"))

    def test_official_dn_range_uses_midpoint_only_as_curve_axis(self):
        interval = dimension_range('TUBO DE AÇO CARBONO, DN 1000 MM (40") OU DN 1100 MM (44"), JUNTA SOLDADA')
        self.assertEqual((interval["low"], interval["high"], interval["x"]), (1000, 1100, 1050))


class RealReferenceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.db = load_database(APP / "data" / "sinapi_base.json.gz")
        cls.families = build_families(cls.db)

    def family(self, code):
        return next(family for family in self.families if any(ref["code"] == code for ref in family["refs"]))

    def test_base_contains_only_scope(self):
        self.assertTrue(self.db["compositions"])
        self.assertTrue(all(is_automatic_scope(item["description"]) for item in self.db["compositions"].values()))
        self.assertNotIn("91924", self.db["compositions"])

    def test_all_dimensional_steel_and_iron_compositions_are_searchable(self):
        import re
        metal = re.compile(r"\b(?:ACO(?: GALVANIZADO| CARBONO| INOXIDAVEL)?|FERRO(?: GALVANIZADO| FUNDIDO)?)\b")
        visible = {ref["code"] for family in self.families for ref in family["refs"]}
        expected = {code for code, item in self.db["compositions"].items()
                    if metal.search(norm(item["description"]))
                    and (connection_geometry(item["description"]) or dimension_range(item["description"]) or measures(item["description"]))}
        self.assertTrue(expected)
        self.assertEqual(expected - visible, set())


    def test_typographic_variants_join_same_family(self):
        family = self.family("89867")
        self.assertEqual({ref["code"] for ref in family["refs"]}, {"104318", "89867", "104320"})
        self.assertEqual([ref["dims"][0] for ref in family["refs"]], [20, 25, 32])

    def test_justification_lists_every_reference_used(self):
        family = self.family("94464")
        result = calculate(self.db, family, "94464", 70)
        used = [ref["code"] for ref in result["references"] if ref["used"]]
        for code in used:
            self.assertIn(code, result["justification"])
        self.assertIn("Códigos de referência:", result["justification"])

    def test_94464_galvanized_steel_family_and_official_target(self):
        family = self.family("94464")
        self.assertEqual({ref["code"] for ref in family["refs"]}, {"94462", "94463", "94464"})
        self.assertEqual([ref["dims"][0] for ref in family["refs"]], [50, 65, 80])
        result = calculate(self.db, family, "94464", 65)
        self.assertEqual(result["target_code"], "94463")
        material = next(row for row in result["rows"] if row["source_resource"] == "7694")
        self.assertEqual(material["target_resource"], "7701")
        self.assertIn('2.1/2"', material["target_description"])

    def test_cast_iron_bare_dn_and_carbon_steel_ranges(self):
        cast_iron = self.family("105379")
        self.assertGreaterEqual(len(cast_iron["refs"]), 10)
        self.assertEqual(next(ref["dims"] for ref in cast_iron["refs"] if ref["code"] == "105379"), [150])
        carbon = self.family("97177")
        first = next(ref for ref in carbon["refs"] if ref["code"] == "97177")
        self.assertEqual(format_reference_dimensions(first), "1000–1100 mm")
        direct = calculate(self.db, carbon, "97177", 1100)
        self.assertEqual(direct["target_code"], "97177")

    def test_normal_and_reducing_tees_are_separate_three_port_families(self):
        reduction = self.family("89624")
        normal = self.family("89617")
        self.assertEqual(reduction["dimension_count"], 3)
        self.assertIn("TE DE REDUCAO", reduction["title"])
        refs = {ref["code"]: ref["dims"] for ref in reduction["refs"]}
        self.assertEqual(refs["89624"], [40, 40, 32])
        self.assertTrue({"89622", "89624", "89626", "89627", "89630", "89632"} <= set(refs))
        self.assertNotIn("89617", refs)
        self.assertNotIn("REDUCAO", normal["title"])
        self.assertNotIn("89624", {ref["code"] for ref in normal["refs"]})

    def test_all_reducing_sleeve_dimensions_share_technical_family(self):
        family = self.family("92928")
        self.assertEqual(family["dimension_count"], 2)
        self.assertIn("LUVA DE REDUCAO", family["title"])
        refs = {ref["code"]: ref["dims"] for ref in family["refs"]}
        self.assertEqual(refs["92928"], [1.5, 1.25])
        self.assertEqual(refs["101932"], [4.0, 3.0])
        self.assertGreaterEqual(len(refs), 18)

    def test_brass_connector_family_uses_both_dimensions(self):
        family = self.family("94607")
        self.assertEqual(family["dimension_count"], 2)
        self.assertEqual(family["axis_units"], ["MM", "POL"])
        self.assertEqual({r["code"] for r in family["refs"]}, {"94607", "94609", "94611", "94613"})
        geometry = connection_geometry(self.db["compositions"]["94609"]["description"])
        self.assertEqual(geometry["dims"], [66, 2.5])
        self.assertEqual(format_geometry(geometry["dims"], geometry["units"]), '66 mm × 2 1/2"')

    def test_all_three_target_dimensions_change_result_and_justification(self):
        family = self.family("89624")
        a = calculate(self.db, family, "89624", [45, 45, 35])
        b = calculate(self.db, family, "89624", [45, 45, 40])
        labor_a = next(row for row in a["rows"] if row["source_resource"] == "88267")
        labor_b = next(row for row in b["rows"] if row["source_resource"] == "88267")
        self.assertNotEqual(labor_a["coefficient"], labor_b["coefficient"])
        self.assertIn("40 mm × 40 mm × 32 mm", a["justification"])
        self.assertIn("45 mm × 45 mm × 35 mm", a["justification"])

    def test_exact_tee_finds_official_input_and_pe_sp_prices(self):
        result = calculate(self.db, self.family("89624"), "89624", [50, 50, 40])
        self.assertEqual(result["target_code"], "89626")
        material = next(row for row in result["rows"] if row["target_resource"] == "7131")
        self.assertIn("PE e SP", material["supply_status"])
        self.assertIsNotNone(material["supply_matches"][0]["PE"])
        self.assertIsNotNone(material["supply_matches"][0]["SP"])

    def test_nonexistent_target_reports_no_input(self):
        result = calculate(self.db, self.family("89624"), "89624", [45, 45, 35])
        material = next(row for row in result["rows"] if row["source_resource"] == "7128")
        self.assertIn("não encontrado", material["supply_status"])

    def test_89355_target_40_uses_new_input_description_and_code(self):
        result = calculate(self.db, self.family("89355"), "89355", "40 mm", True)
        material = next(row for row in result["rows"] if row["source_resource"] == "9867")
        self.assertIn("20 MM", material["description"])
        self.assertEqual(material["target_resource"], "9874")
        self.assertIn("40 MM", material["target_description"])
        self.assertTrue(result['justification'].startswith('Banco: SINAPI. Data do banco: 08/2026.'))
        self.assertIn('89355 (20 mm)', result['justification'])
        self.assertIn("de 20 mm para 40 mm", result["justification"])

    def test_report_has_justification_prices_and_formula_safety(self):
        from openpyxl import load_workbook
        result = calculate(self.db, self.family("89624"), "89624", [50, 50, 40])
        result["rows"][0]["description"] = "=THIS_IS_TEXT()"
        with tempfile.TemporaryDirectory() as folder:
            file = export_report(result, Path(folder) / "relatorio.xlsx")
            workbook = load_workbook(file, data_only=False)
            self.assertEqual(len(workbook.sheetnames), 5)
            self.assertEqual(workbook["Coeficientes"]["C2"].data_type, "s")
            headers = [cell.value for cell in workbook["Coeficientes"][1]]
            self.assertIn("Código do insumo na nova composição", headers)
            self.assertIn("Descrição do insumo na nova composição", headers)
            self.assertIn("Justificativa curta", [cell.value for cell in workbook["Leia primeiro"]["A"]])
            rows = list(workbook["Insumos SINAPI PE-SP"].iter_rows(min_row=2, values_only=True))
            self.assertIn("7131", [row[2] for row in rows])
            workbook.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
