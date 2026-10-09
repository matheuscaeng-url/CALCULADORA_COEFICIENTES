"""Regressões: formato de justificativas e Excel sem bibliotecas instaladas."""
import copy
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP))

from coeficientes_core import build_families, calculate, export_report, load_database
from regra_tres_core import calculate_generic_rule_of_three, calculate_rule_of_three
from tendencia_manual_core import (calculate_generic_parameter_trend, calculate_manual_trend,
                                   series_from_sinapi)


def release_results():
    db = load_database(APP / 'data/sinapi_base.json.gz')
    families = build_families(db)
    def family(code):
        return next(f for f in families if any(r['code'] == code for r in f['refs']))
    trend = calculate(db, family('89400'), '89400', [50, 50, 32], True)
    rule = calculate_rule_of_three(db, family('90734'), 200)
    refs = [{'code': code, 'dims': [d], 'units': ['MM']}
            for code, d in [('89355', 20), ('89356', 25)]]
    series, unit, _ = series_from_sinapi(db, refs)
    manual = calculate_manual_trend(refs, series, [23], service_unit=unit, meta=db['meta'])
    return {'tendencia': trend, 'regra_sinapi': rule, 'manual_sinapi': manual}


class JustificationReleaseTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.results = release_results()

    def check_order(self, text, bank, period):
        self.assertTrue(text.startswith(f'Banco: {bank}. Data do banco: {period}.'))
        reference_index = text.index(' de referência:')
        self.assertLess(reference_index, text.index('Técnica:'))
        self.assertNotIn('Resultados:', text)

    def test_all_five_calculators_have_same_field_order(self):
        for name, result in self.results.items():
            with self.subTest(name=name):
                self.check_order(result['justification'], 'SINAPI', '08/2026')
        rule = calculate_generic_rule_of_three([25, 20], [75, 50], .0844, '1217', 'ORSE', '07/26')
        self.check_order(rule['justification'], 'ORSE', '07/26')
        trend = calculate_generic_parameter_trend(
            [{'code': '1213', 'parameter': '1', 'coefficients': [.16]},
             {'code': '1217', 'parameter': '2', 'coefficients': [.36]}],
            '4', True, 'POL', bank='ORSE', period='07/26')
        self.check_order(trend['justification'], 'ORSE', '07/26')
        self.assertIn('1213 (1")', trend['justification'])
        self.assertIn('1217 (2")', trend['justification'])
        self.assertIn('Extrapolação linear', trend['justification'])
        self.assertAlmostEqual(trend['coefficients'][0], .76)

    def test_rule_of_three_factor_is_last_without_formula_or_derivation(self):
        generic = calculate_generic_rule_of_three([25], [50], .1, 'A', 'Próprio', '10/2026')
        self.assertTrue(generic['justification'].endswith('Fator de correção utilizado: 2.'))
        for result in [generic, self.results['regra_sinapi']]:
            text = result['justification']
            for token in ['=', '×', ' / ', 'soma', 'dimensões', 'RMS']:
                self.assertNotIn(token, text)
            factor = f"{result['factor']:.12g}".replace('.', ',')
            self.assertTrue(text.endswith(f'Fator de correção utilizado: {factor}.'))

    def test_automatic_two_point_extrapolation_is_named_linear(self):
        result = self.results['tendencia']
        self.assertIn('linear', result['method'].lower())
        self.assertIn('linear', result['justification'].lower())
        self.assertNotIn('PCHIP', result['justification'])
        for ref in result['references']:
            if ref['used']:
                self.assertIn(ref['code'], result['justification'])

    def test_exact_reference_does_not_claim_interpolation(self):
        refs = [{'code': 'A', 'parameter': 10, 'coefficients': [.1]},
                {'code': 'B', 'parameter': 20, 'coefficients': [.3]}]
        result = calculate_generic_parameter_trend(refs, 20, bank='ORSE', period='07/26')
        self.assertIn('Código de referência: B', result['justification'])
        self.assertIn('Consulta ao ponto informado', result['justification'])
        self.assertNotIn('Códigos de referência:', result['justification'])


class ExcelReleaseTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.results = release_results()

    def test_all_export_buttons_save_readable_reports(self):
        from openpyxl import load_workbook
        from coeficientes_app import CalculatorApp
        from calculadoras_extras import OutputPanel
        for name, result in self.results.items():
            with self.subTest(calculator=name), tempfile.TemporaryDirectory() as folder:
                destination = Path(folder) / ('Relatório_' + name + '.xlsx')
                app = Mock(working=False)
                app.start_job.side_effect = lambda work, done, message: done(work())
                if name == 'tendencia':
                    app.result = result
                    with patch('coeficientes_app.filedialog.asksaveasfilename', return_value=str(destination)), \
                         patch('coeficientes_app.messagebox.showinfo'):
                        CalculatorApp.export(app)
                else:
                    panel = Mock(result=result, app=app)
                    with patch('calculadoras_extras.filedialog.asksaveasfilename', return_value=str(destination)):
                        OutputPanel.export(panel)
                self.assertTrue(destination.is_file())
                book = load_workbook(destination)
                self.assertEqual(len(book.sheetnames), 5)
                metadata = dict(book['Leia primeiro'].iter_rows(values_only=True))
                self.assertEqual(metadata['Banco'], 'SINAPI')
                self.assertEqual(metadata['Competência'], '08/2026')
                self.assertEqual(metadata['Justificativa curta'], result['justification'])
                rows = list(book['Coeficientes'].iter_rows(min_row=2, values_only=True))
                self.assertEqual(len(rows), len(result['rows']))
                for record, source in zip(rows, result['rows']):
                    self.assertAlmostEqual(record[7], source['coefficient'])
                book.close()

    def test_exports_work_without_site_packages_or_installed_openpyxl(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'resultados.json'
            source.write_text(json.dumps(self.results), encoding='utf-8')
            script = '''import importlib.util, json, sys
from pathlib import Path
assert importlib.util.find_spec('openpyxl') is None
from coeficientes_core import export_report
import openpyxl, et_xmlfile
assert 'dependencias' in openpyxl.__file__
assert 'dependencias' in et_xmlfile.__file__
source = Path(sys.argv[1])
for name, result in json.loads(source.read_text(encoding='utf-8')).items():
    target = export_report(result, source.parent / (name + '.xlsx'))
    book = openpyxl.load_workbook(target)
    assert len(book.sheetnames) == 5
    book.close()
print('3 exports OK without site-packages')
'''
            env = dict(os.environ)
            env.pop('PYTHONPATH', None)
            completed = subprocess.run([sys.executable, '-S', '-c', script, str(source)],
                cwd=APP, env=env, text=True, capture_output=True, timeout=45)
            self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
            self.assertIn('3 exports OK', completed.stdout)


if __name__ == '__main__':
    unittest.main()
